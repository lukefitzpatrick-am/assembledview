#!/usr/bin/env python3
"""Build a v5 campaign report from content.json. Numbers come from report-data.json keys.

usage: build_report.py content.json report-data.json template.potx out.pptx photo1 photo2 photo3
"""
import json
import os
import re
import sys

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches
from lxml import etree

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from am_rebrand import (  # noqa: E402
    BODY, HEAD, K, Deck, I, Run,
    arch, btn, card, journey, line, numdot, pill, table, tag, tb, trio,
)

KEY_RE = re.compile(r"[A-Za-z][A-Za-z0-9_]*(\.[A-Za-z0-9_]+)+")
SERIF_RE = re.compile(r"\{i:\"([^\"]*)\"\}")
CHANNEL_KEY_RE = re.compile(r"channels\.(\d+)\..+")
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
AXIS_FORMAT = {"money": "$#,##0", "count": "#,##0", "percent": "0.00%"}
CONVERSION_NOTE = "Conversion rate is results divided by clicks across live channels."
FIXED_STATS = (
    ("Spend", "spend"),
    ("Impressions", "impressions"),
    ("Clicks", "clicks"),
    ("CTR", "metrics.ctr"),
)
DOTS = (
    (K["sky"], K["ink"]),
    (K["forest"], K["white"]),
    (K["lime"], K["ink"]),
    (K["forest"], K["white"]),
)


class BuildError(Exception):
    pass


def fail(path, message):
    raise BuildError(f"{path}: {message}")


def slug(value):
    return re.sub(r"[^a-z0-9]+", "-", str(value).lower()).strip("-")


def matches(item, part):
    group = str(item.get("group") or "")
    label = slug(item.get("label") or "")
    return part in (group, label, label.split("-")[0], group.split("_")[0])


def resolve(data, key, at):
    if not isinstance(key, str) or not key.strip():
        fail(at, f"missing key {key!r}")
    current = data
    for part in key.split("."):
        if isinstance(current, list):
            if part.isdigit():
                index = int(part)
                if index < 0 or index >= len(current):
                    fail(at, f"missing key {key}")
                current = current[index]
                continue
            found = next((item for item in current if isinstance(item, dict) and matches(item, part)), None)
            if found is None:
                fail(at, f"missing key {key}")
            current = found
            continue
        if isinstance(current, dict) and part in current:
            current = current[part]
            continue
        fail(at, f"missing key {key}")
    if not isinstance(current, str):
        fail(at, f"missing key {key}")
    return current


def digits(text):
    return re.search(r"\d", text) is not None


def text_of(node, path, limit=None, words=None):
    if isinstance(node, str) and KEY_RE.fullmatch(node):
        out = resolve(DATA, node, path)
    elif isinstance(node, str):
        if digits(node):
            fail(path, "digits in copy; use a value key")
        out = node
    elif isinstance(node, list):
        bits = []
        for index, part in enumerate(node):
            at = f"{path}.{index}"
            if isinstance(part, str):
                if digits(part):
                    fail(at, "digits in copy; use a value key")
                bits.append(part)
            elif isinstance(part, dict) and isinstance(part.get("k"), str) and len(part) == 1:
                bits.append(resolve(DATA, part["k"], f"{at}.k"))
            else:
                fail(at, 'expected copy or {"k": path}')
        out = "".join(bits)
    else:
        fail(path, "expected copy or a value key")
    if limit is not None and len(out) > limit:
        fail(path, f"{len(out)} characters, limit {limit}")
    if words is not None and len(re.findall(r"\S+", out)) > words:
        fail(path, f"over {words} words")
    return out


def key_of(node, path):
    if not isinstance(node, str):
        fail(path, "expected a value key")
    return resolve(DATA, node, path)


def parse_title(title, path):
    if not isinstance(title, str):
        fail(path, "title must be a string")
    if len(SERIF_RE.findall(title)) > 1:
        fail(path, "one serif phrase")
    visible = SERIF_RE.sub(lambda match: match.group(1), title)
    if digits(visible):
        fail(path, "digits in copy; use a value key")
    if len(visible) > 60:
        fail(path, f"{len(visible)} characters, limit 60")
    if not visible.endswith("."):
        fail(path, "title must end with a full stop")
    parts = []
    cursor = 0
    for match in SERIF_RE.finditer(title):
        if match.start() > cursor:
            parts.append(title[cursor:match.start()])
        phrase = match.group(1)
        cursor = match.end()
        # A stop that ends the title here belongs on this serif run, not in a run of its own.
        stop = re.match(r"[ \t]*\.", title[cursor:])
        if stop and cursor + stop.end() == len(title):
            phrase = phrase.rstrip() + "."
            cursor += stop.end()
        parts.append(I(phrase))
    if cursor < len(title):
        parts.append(re.sub(r"[ \t]+([.!?])$", r"\1", title[cursor:]))
    if len(parts) >= 2 and isinstance(parts[-1], str) and re.fullmatch(r"[ \t]*[.!?]", parts[-1]):
        punct = parts[-1].strip()
        prev = parts[-2]
        if isinstance(prev, Run):
            parts[-2] = Run(prev.t.rstrip() + punct, **prev.k)
        else:
            parts[-2] = prev.rstrip() + punct
        parts.pop()
    parts = [part for part in parts if part != ""]
    return visible, parts or [visible]


def no_autofit(frame):
    frame.word_wrap = True
    body = frame._txBody.find(qn("a:bodyPr"))
    if body is None:
        return
    for tag in ("a:normAutofit", "a:spAutoFit"):
        for node in body.findall(qn(tag)):
            body.remove(node)
    if body.find(qn("a:noAutofit")) is None:
        etree.SubElement(body, qn("a:noAutofit"))


def put_title(deck, slide, index, spec, path, size=None, anchor="t"):
    visible, runs = parse_title(spec["title"], path)
    if size is None:
        size = 32 if len(visible) > 36 else 36
    color = K["white"] if slide._am_dark else K["ink"]
    shape = deck.fill(slide, index, runs, base=dict(size=size, color=color, font=HEAD), align="l")
    no_autofit(shape.text_frame)
    shape.text_frame.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    shape.name = path
    return shape


def say(slide, x, y, w, h, content, path, size=13, color=None, font=None, align=None, anchor="t"):
    shape = tb(
        slide, x, y, w, h, content, size=size, color=color or K["ink"],
        font=font or BODY, align=align, valign=anchor,
    )
    shape.name = path
    no_autofit(shape.text_frame)
    return shape


def fill_named(deck, slide, index, content, path, size, color, font=None):
    shape = deck.fill(slide, index, content, base=dict(size=size, color=color, font=font or BODY), align="l")
    no_autofit(shape.text_frame)
    shape.name = path
    return shape


def require(spec, field, path):
    if field not in spec:
        fail(path, f"missing {field}")
    return spec[field]


def variant(spec, path, allowed):
    name = spec.get("variant")
    if name not in allowed:
        fail(path, f"variant must be {' or '.join(allowed)}")
    return name


def with_pace(label):
    stripped = label.strip()
    if stripped.lower().rstrip(".") == "spend pace":
        return stripped.rstrip(".") + " of expected to date"
    return label


def kpi_figure(label, value):
    """Drop a trailing unit that repeats the KPI label. Frequency + '1.0 frequency' is '1.0'."""
    text = value.strip()
    units = []
    cleaned = label.strip()
    if cleaned:
        units.append(cleaned)
    units.extend(re.findall(r"[A-Za-z]{2,}", cleaned))
    seen = set()
    for unit in units:
        key = unit.casefold()
        if key in seen:
            continue
        seen.add(key)
        trimmed = re.sub(rf"\s+{re.escape(unit)}\s*$", "", text, count=1, flags=re.IGNORECASE)
        if trimmed != text and trimmed.strip():
            text = trimmed.strip()
    return text


def fitted_copy_height(text, width, size):
    """0.3in pad, one line-height per wrapped line, 0.3in pad. At least 0.9in."""
    line_h = size * 1.15 / 72
    lines = wrap_lines(text, width, size / 72 * 0.55)
    return max(0.9, 0.3 + lines * line_h + 0.3)


def insight_card(slide, path, spec, above_bottom, height=None):
    """Sand insight card, 0.3in below the element above it. Height fits the text."""
    insight = text_of(require(spec, "insight", f"{path}.insight"), f"{path}.insight", 160)
    top = above_bottom + 0.3
    text_w = 10.1
    if height is None:
        height = fitted_copy_height(insight, text_w, 15)
    card(slide, 0.6, top, 12.13, height, K["sand"])
    tag(slide, 0.82, top + (height / 2) - 0.16, "Insight", K["sky"], K["ink"])
    shape = say(
        slide, 2.35, top + 0.3, text_w, height - 0.6, insight, f"{path}.insight",
        size=15, color=K["ink"], anchor="t",
    )
    shape.text_frame.margin_top = 0
    shape.text_frame.margin_bottom = 0
    return top + height


def cq(tag):
    return f"{{{C_NS}}}{tag}"


def aq(tag):
    return f"{{{A_NS}}}{tag}"


def lookup(data, key):
    """Walk a dotted path and return the value, or None when it is missing."""
    if not isinstance(data, dict) or not isinstance(key, str) or not key.strip():
        return None
    current = data
    for part in key.split("."):
        if isinstance(current, list):
            if part.isdigit():
                index = int(part)
                if index < 0 or index >= len(current):
                    return None
                current = current[index]
                continue
            found = next((item for item in current if isinstance(item, dict) and matches(item, part)), None)
            if found is None:
                return None
            current = found
            continue
        if isinstance(current, dict) and part in current:
            current = current[part]
            continue
        return None
    return current


def as_chart(data, key):
    found = lookup(data, key)
    if not isinstance(found, dict):
        return None
    points = found.get("points")
    if not isinstance(points, list) or not points:
        return None
    return found


def channel_index_for_group(data, group):
    channels = data.get("channels") if isinstance(data, dict) else None
    if not isinstance(channels, list) or not isinstance(group, str) or not group:
        return None
    for index, channel in enumerate(channels):
        if isinstance(channel, dict) and channel.get("group") == group:
            return index
    return None


def named_channel_group(spec, data):
    """The payload group named by a slide channel field, when it names one."""
    raw = spec.get("channel")
    channels = data.get("channels") if isinstance(data, dict) else None
    if not isinstance(raw, str) or not isinstance(channels, list):
        return None
    match = CHANNEL_KEY_RE.fullmatch(raw)
    if match:
        index = int(match.group(1))
        if 0 <= index < len(channels) and isinstance(channels[index], dict):
            group = channels[index].get("group")
            return group if isinstance(group, str) else None
    for channel in channels:
        if isinstance(channel, dict) and raw in (channel.get("group"), channel.get("label")):
            group = channel.get("group")
            return group if isinstance(group, str) else None
    return None


def channel_chart_problem(spec, data):
    """None when the chart key is this slide's channel chart. Otherwise the fix."""
    if not isinstance(data, dict):
        return "use variant cards"
    chart = as_chart(data, spec.get("chart"))
    if chart is None or chart.get("scope") != "channel":
        return "use variant cards"
    group = chart.get("group")
    if channel_index_for_group(data, group) is None:
        return "use variant cards"
    expected = named_channel_group(spec, data)
    if expected is not None and expected != group:
        return "chart group must match this channel. Point chart at that channel's chart, or use variant cards."
    return None


def chart_txpr(size="1000", wrap="square"):
    tx_pr = etree.Element(cq("txPr"))
    body = etree.SubElement(tx_pr, aq("bodyPr"))
    body.set("wrap", wrap)
    etree.SubElement(tx_pr, aq("lstStyle"))
    paragraph = etree.SubElement(tx_pr, aq("p"))
    props = etree.SubElement(paragraph, aq("pPr"))
    run = etree.SubElement(props, aq("defRPr"))
    run.set("sz", size)
    solid = etree.SubElement(run, aq("solidFill"))
    etree.SubElement(solid, aq("srgbClr")).set("val", "3B4A3F")
    for face in ("latin", "ea", "cs"):
        etree.SubElement(run, aq(face)).set("typeface", "Plus Jakarta Sans")
    etree.SubElement(paragraph, aq("endParaRPr")).set("sz", size)
    return tx_pr


def put_txpr(parent, size="1000", wrap="square"):
    current = parent.find(cq("txPr"))
    fresh = chart_txpr(size, wrap)
    if current is not None:
        parent.replace(current, fresh)
        return
    cross = parent.find(cq("crossAx"))
    if cross is not None:
        cross.addprevious(fresh)
    else:
        parent.append(fresh)


def set_num_fmt(axis, code):
    node = axis.find(cq("numFmt"))
    if node is None:
        node = etree.Element(cq("numFmt"))
        tick = axis.find(cq("majorTickMark"))
        if tick is not None:
            tick.addprevious(node)
        else:
            axis.append(node)
    node.set("formatCode", code)
    node.set("sourceLinked", "0")


def style_major_gridlines(axis):
    grid = axis.find(cq("majorGridlines"))
    if grid is None:
        grid = etree.Element(cq("majorGridlines"))
        pos = axis.find(cq("axPos"))
        if pos is not None:
            pos.addnext(grid)
        else:
            axis.append(grid)
    for child in list(grid):
        grid.remove(child)
    props = etree.SubElement(grid, cq("spPr"))
    stroke = etree.SubElement(props, aq("ln"))
    stroke.set("w", "9525")
    solid = etree.SubElement(stroke, aq("solidFill"))
    etree.SubElement(solid, aq("srgbClr")).set("val", "E4DED2")


def series_solid(ser, colour):
    props = ser.find(cq("spPr"))
    if props is not None:
        ser.remove(props)
    props = etree.Element(cq("spPr"))
    solid = etree.SubElement(props, aq("solidFill"))
    etree.SubElement(solid, aq("srgbClr")).set("val", colour)
    stroke = etree.SubElement(props, aq("ln"))
    etree.SubElement(stroke, aq("noFill"))
    title = ser.find(cq("tx"))
    if title is not None:
        title.addnext(props)
    else:
        ser.append(props)


def series_line(ser):
    props = ser.find(cq("spPr"))
    if props is not None:
        ser.remove(props)
    props = etree.Element(cq("spPr"))
    stroke = etree.SubElement(props, aq("ln"))
    stroke.set("w", "25400")
    solid = etree.SubElement(stroke, aq("solidFill"))
    etree.SubElement(solid, aq("srgbClr")).set("val", "0F1D13")
    title = ser.find(cq("tx"))
    if title is not None:
        title.addnext(props)
    else:
        ser.append(props)
    marker = ser.find(cq("marker"))
    if marker is not None:
        ser.remove(marker)
    marker = etree.Element(cq("marker"))
    etree.SubElement(marker, cq("symbol")).set("val", "none")
    props.addnext(marker)
    smooth = ser.find(cq("smooth"))
    if smooth is not None:
        ser.remove(smooth)
    smooth = etree.Element(cq("smooth"))
    smooth.set("val", "0")
    ser.append(smooth)


def no_chart_border(space):
    props = space.find(cq("spPr"))
    if props is not None:
        space.remove(props)
    props = etree.Element(cq("spPr"))
    stroke = etree.SubElement(props, aq("ln"))
    etree.SubElement(stroke, aq("noFill"))
    chart = space.find(cq("chart"))
    if chart is not None:
        chart.addnext(props)
    else:
        space.append(props)


def add_secondary_line(chart):
    """Move the second column series onto a line chart with its own value axis.

    python-pptx has no combo chart. The line series stays in the embedded
    workbook, so the client can still edit both series.
    """
    space = chart._element
    plot = space.find(cq("chart")).find(cq("plotArea"))
    bars = plot.find(cq("barChart"))
    if bars is None:
        return
    series = bars.findall(cq("ser"))
    if len(series) < 2:
        return
    line_ser = series[-1]
    bars.remove(line_ser)
    series_line(line_ser)
    cat_ax = plot.find(cq("catAx"))
    val_ax = plot.find(cq("valAx"))
    cat_id = cat_ax.find(cq("axId")).get("val")
    primary_id = val_ax.find(cq("axId")).get("val")
    try:
        secondary_id = str(int(primary_id) + 1)
    except (TypeError, ValueError):
        secondary_id = "100"
    if secondary_id in (cat_id, primary_id):
        secondary_id = "100"
    line = etree.Element(cq("lineChart"))
    etree.SubElement(line, cq("grouping")).set("val", "standard")
    etree.SubElement(line, cq("varyColors")).set("val", "0")
    line.append(line_ser)
    etree.SubElement(line, cq("marker")).set("val", "0")
    etree.SubElement(line, cq("smooth")).set("val", "0")
    etree.SubElement(line, cq("axId")).set("val", cat_id)
    etree.SubElement(line, cq("axId")).set("val", secondary_id)
    bars.addnext(line)
    secondary = etree.fromstring(etree.tostring(val_ax))
    secondary.find(cq("axId")).set("val", secondary_id)
    secondary.find(cq("axPos")).set("val", "r")
    grid = secondary.find(cq("majorGridlines"))
    if grid is not None:
        secondary.remove(grid)
    crosses = secondary.find(cq("crosses"))
    if crosses is None:
        crosses = etree.SubElement(secondary, cq("crosses"))
    crosses.set("val", "max")
    cross = secondary.find(cq("crossAx"))
    if cross is not None:
        cross.set("val", cat_id)
    val_ax.addnext(secondary)


def date_label(value):
    text = str(value)
    match = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", text)
    if not match:
        return text
    month = int(match.group(2))
    if month < 1 or month > 12:
        return text
    return f"{int(match.group(3))} {MONTHS[month - 1]}"


def axis_format(axis):
    if not isinstance(axis, dict):
        return AXIS_FORMAT["count"]
    return AXIS_FORMAT.get(axis.get("format"), AXIS_FORMAT["count"])


def point_value(point, index):
    if not isinstance(point, (list, tuple)) or len(point) <= index or point[index] is None:
        return None
    return point[index]


def manual_layout(x, y, w, h):
    layout = etree.Element(cq("layout"))
    manual = etree.SubElement(layout, cq("manualLayout"))
    for tag, val in (
        ("xMode", "edge"),
        ("yMode", "edge"),
        ("x", x),
        ("y", y),
        ("w", w),
        ("h", h),
    ):
        etree.SubElement(manual, cq(tag)).set("val", val)
    return layout


CHART_TOP = 2.85
CONTENT_FLOOR = 6.7
CHART_NOTE_H = 0.22
INSIGHT_GAP = 0.3
MIN_CHART_H = 1.2


def chart_slide_heights(chart, insight):
    """Chart, note and insight stay inside y 2.85 to 6.7. Shrink the chart first."""
    note = chart.get("note") if isinstance(chart, dict) else None
    note_h = CHART_NOTE_H if isinstance(note, str) and note.strip() else 0.0
    insight_h = fitted_copy_height(insight, 10.1, 15)
    chart_h = CONTENT_FLOOR - CHART_TOP - note_h - INSIGHT_GAP - insight_h
    if chart_h < MIN_CHART_H:
        chart_h = MIN_CHART_H
        insight_h = CONTENT_FLOOR - CHART_TOP - chart_h - note_h - INSIGHT_GAP
    return chart_h, insight_h


def place_legend(space):
    """Legend is the top 8%, top right. The plot fills the rest, inset for the axes and dates."""
    node = space.find(cq("chart"))
    legend = node.find(cq("legend"))
    if legend is None:
        return
    pos = legend.find(cq("legendPos"))
    if pos is None:
        pos = etree.Element(cq("legendPos"))
        legend.insert(0, pos)
    pos.set("val", "tr")
    overlay = legend.find(cq("overlay"))
    if overlay is None:
        overlay = etree.SubElement(legend, cq("overlay"))
    overlay.set("val", "0")
    current = legend.find(cq("layout"))
    if current is not None:
        legend.remove(current)
    overlay.addprevious(manual_layout("0.50", "0.00", "0.50", "0.08"))
    put_txpr(legend, "900")
    plot = node.find(cq("plotArea"))
    current = plot.find(cq("layout"))
    if current is not None:
        plot.remove(current)
    plot.insert(0, manual_layout("0.08", "0.10", "0.84", "0.74"))


def style_delivery_chart(chart, left_fmt, right_fmt, skip):
    space = chart._element
    node = space.find(cq("chart"))
    title = node.find(cq("title"))
    if title is not None:
        node.remove(title)
    auto = node.find(cq("autoTitleDeleted"))
    if auto is None:
        auto = etree.Element(cq("autoTitleDeleted"))
        node.insert(0, auto)
    auto.set("val", "1")
    plot = node.find(cq("plotArea"))
    bars = plot.find(cq("barChart"))
    if bars is not None:
        vary = bars.find(cq("varyColors"))
        if vary is None:
            vary = etree.Element(cq("varyColors"))
            grouping = bars.find(cq("grouping"))
            if grouping is not None:
                grouping.addnext(vary)
            else:
                bars.insert(0, vary)
        vary.set("val", "0")
        series = bars.findall(cq("ser"))
        if series:
            series_solid(series[0], "246646")
    cat_ax = plot.find(cq("catAx"))
    skip_node = cat_ax.find(cq("tickLblSkip"))
    if skip_node is None:
        skip_node = etree.Element(cq("tickLblSkip"))
        anchor = cat_ax.find(cq("noMultiLvlLbl"))
        if anchor is not None:
            anchor.addprevious(skip_node)
        else:
            cat_ax.append(skip_node)
    skip_node.set("val", str(skip))
    auto = cat_ax.find(cq("auto"))
    if auto is None:
        auto = etree.Element(cq("auto"))
        skip_node.addprevious(auto)
    auto.set("val", "0")
    put_txpr(cat_ax, wrap="none")
    for axis in plot.findall(cq("valAx")):
        pos = axis.find(cq("axPos"))
        side = pos.get("val") if pos is not None else "l"
        if side == "l":
            set_num_fmt(axis, left_fmt)
            style_major_gridlines(axis)
        else:
            set_num_fmt(axis, right_fmt)
            grid = axis.find(cq("majorGridlines"))
            if grid is not None:
                axis.remove(grid)
        put_txpr(axis)
    put_txpr(space)
    no_chart_border(space)
    place_legend(space)
    area = plot.find(cq("spPr"))
    if area is not None:
        plot.remove(area)
    area = etree.Element(cq("spPr"))
    stroke = etree.SubElement(area, aq("ln"))
    etree.SubElement(stroke, aq("noFill"))
    plot.append(area)


def delivery_chart(slide, x, y, w, h, chart, path):
    """Native column chart, with the right series as a line on a second axis."""
    points = chart.get("points") if isinstance(chart, dict) else None
    if not isinstance(points, list) or not points:
        fail(path, "chart must resolve to an object with points. Point chart at charts.0.")
    categories = []
    left_values = []
    right_values = []
    right = chart.get("right") if isinstance(chart.get("right"), dict) else None
    for point in points:
        categories.append(date_label(point[0] if isinstance(point, (list, tuple)) and point else ""))
        left = point_value(point, 1)
        left_values.append(0 if left is None else left)
        if right is not None:
            right_values.append(point_value(point, 2))
    data = CategoryChartData()
    data.categories = categories
    left_axis = chart.get("left") if isinstance(chart.get("left"), dict) else {}
    data.add_series(str(left_axis.get("label") or "Value"), left_values)
    if right is not None:
        data.add_series(str(right.get("label") or "Value"), right_values)
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data,
    )
    frame.name = f"{path}.chart"
    graphic = frame.chart
    graphic.has_legend = True
    graphic.legend.position = XL_LEGEND_POSITION.CORNER
    graphic.legend.include_in_layout = True
    graphic.plots[0].gap_width = 60
    add_secondary_line(graphic)
    count = len(points)
    skip = max(1, (count + 7) // 8)
    style_delivery_chart(graphic, axis_format(left_axis), axis_format(right), skip)
    bottom = y + h
    note = chart.get("note")
    if isinstance(note, str) and note.strip():
        say(slide, x, bottom, w, 0.22, note, f"{path}.note", size=9, color=K["muted"])
        bottom += 0.22
    return bottom


def stat_pills(slide, path, entries):
    """entries are (label, value, shape name). One row, y 1.95, height 0.75."""
    if not entries:
        return
    gap = 0.14
    width = (12.13 - gap * (len(entries) - 1)) / len(entries)
    for index, (label, value, at) in enumerate(entries):
        x = 0.6 + index * (width + gap)
        pill(slide, x, 1.95, width, 0.75, K["sand"])
        say(slide, x + 0.14, 2.02, width - 0.28, 0.34, value, at, size=14, color=K["forest"], font=HEAD)
        say(slide, x + 0.14, 2.34, width - 0.28, 0.26, label, f"{at}.label", size=11, color=K["ink"])


def wrap_lines(text, width, per_inch):
    chars = max(8, int(max(width, 0.8) / per_inch))
    return max(1, (len(text) + chars - 1) // chars)


def fitted_card_height(title, body, width):
    """Top pad, dot, title, body, tag, bottom pad. Clamped to 2.2–3.4in."""
    inner = max(1.0, width - 0.5)
    title_lines = wrap_lines(title, inner, 0.095)
    body_lines = wrap_lines(body, inner, 0.078)
    height = 0.25 + 0.42 + 0.08 + (0.32 * title_lines) + 0.06 + (0.24 * body_lines) + 0.08 + 0.36 + 0.25
    return max(2.2, min(3.4, height))


def draw_fit_cards(slide, cards, top, left, right):
    """cards: title, body, tag, fill, ink, tag_fill, tag_ink, title_path, body_path."""
    gap = 0.16
    count = len(cards)
    width = (right - left - gap * (count - 1)) / count
    height = max(fitted_card_height(item["title"], item["body"], width) for item in cards)
    for index, item in enumerate(cards):
        x = left + index * (width + gap)
        card(slide, x, top, width, height, K["sand"])
        y = top + 0.25
        numdot(slide, x + 0.18, y, index + 1, item["fill"], item["ink"], 0.36)
        y += 0.44
        title_h = 0.32 * wrap_lines(item["title"], width - 0.5, 0.095)
        say(
            slide, x + 0.18, y, width - 0.36, title_h, item["title"], item["title_path"],
            size=14, color=item.get("title_color", K["ink"]), font=HEAD,
        )
        y += title_h + 0.06
        tag_y = top + height - 0.25 - 0.34
        say(
            slide, x + 0.18, y, width - 0.36, max(0.36, tag_y - 0.08 - y), item["body"], item["body_path"],
            size=13, color=K["body"],
        )
        tag(slide, x + 0.18, tag_y, item["tag"], item["tag_fill"], item["tag_ink"])
    return height


def fixed_stat_values(path, index):
    entries = []
    for label, field in FIXED_STATS:
        key = f"channels.{index}.{field}"
        entries.append((label, key_of(key, f"{path}.{field}"), f"{path}.{field}"))
    return entries


def channel_index_from_pointer(pointer, path):
    if not isinstance(pointer, str):
        fail(path, "expected a channels.<i> key, such as channels.0.label")
    match = CHANNEL_KEY_RE.fullmatch(pointer)
    if not match:
        fail(path, "expected a channels.<i> key, such as channels.0.label")
    return int(match.group(1))


def build_delivery_layout(deck, spec, index, fixed):
    path = f"slides.{index}"
    slide = deck.add("Title Only - White")
    put_title(deck, slide, 100, spec, f"{path}.title")
    if fixed:
        problem = channel_chart_problem(spec, DATA)
        if problem:
            fail(path if problem == "use variant cards" else f"{path}.chart", problem)
        chart = as_chart(DATA, spec.get("chart"))
        channel_index = channel_index_for_group(DATA, chart.get("group"))
        stat_pills(slide, path, fixed_stat_values(path, channel_index))
    else:
        chart_key = require(spec, "chart", f"{path}.chart")
        chart = as_chart(DATA, chart_key)
        if chart is None:
            fail(f"{path}.chart", "chart must resolve to an object with points. Point chart at charts.0.")
        stats = require(spec, "stats", f"{path}.stats")
        if not isinstance(stats, list) or len(stats) > 4:
            fail(f"{path}.stats", "up to 4 stats")
        entries = []
        for stat_index, stat in enumerate(stats):
            at = f"{path}.stats.{stat_index}"
            if not isinstance(stat, dict):
                fail(at, "expected an object")
            label = with_pace(text_of(require(stat, "label", f"{at}.label"), f"{at}.label", 24))
            value = key_of(require(stat, "key", f"{at}.key"), f"{at}.key")
            entries.append((label, value, f"{at}.key"))
        stat_pills(slide, path, entries)
    insight = text_of(require(spec, "insight", f"{path}.insight"), f"{path}.insight", 160)
    chart_h, insight_h = chart_slide_heights(chart, insight)
    chart_bottom = delivery_chart(slide, 0.6, CHART_TOP, 12.13, chart_h, chart, path)
    insight_card(slide, path, spec, chart_bottom, insight_h)
    return slide


def fg(slide):
    return K["white"] if slide._am_dark else K["ink"]


def muted(slide):
    return K["mutedB"] if slide._am_dark else K["muted"]


def magnitude(text):
    raw = text.replace("$", "").replace(",", "").replace("%", "").strip()
    try:
        return abs(float(raw))
    except ValueError:
        return 0.0


def photo_at(photos, index, path):
    if not isinstance(index, int) or isinstance(index, bool) or index < 0 or index >= len(photos):
        fail(path, "photo index out of range")
    return photos[index]


def build_cover(deck, spec, index, photos):
    path = f"slides.{index}"
    kind = variant(spec, path, ("sand", "black"))
    slide = deck.clone(3 if kind == "sand" else 2)
    color = K["white"] if kind == "black" else K["ink"]
    visible, runs = parse_title(require(spec, "title", f"{path}.title"), f"{path}.title")
    size = 32 if len(visible) > 28 else 36
    shape = deck.fill(slide, 103, runs, base=dict(size=size, color=color, font=HEAD), align="l")
    no_autofit(shape.text_frame)
    shape.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    shape.name = f"{path}.title"
    subtitle = key_of(require(spec, "subtitle", f"{path}.subtitle"), f"{path}.subtitle")
    sub = deck.fill(slide, 104, subtitle, base=dict(size=16, color=muted(slide), font=BODY), align="l")
    no_autofit(sub.text_frame)
    sub.name = f"{path}.subtitle"
    for slot, photo in zip((105, 106, 107), photos):
        deck.photo(slide, slot, photo, kind="arch")
    # Placeholder 102 is the client logo slot. With no logo file, leave it off the cover.
    deck.drop(slide, 101, 102)
    return slide


def build_summary(deck, spec, index):
    path = f"slides.{index}"
    kind = variant(spec, path, ("stat-row", "statement"))
    if kind == "statement":
        slide = deck.add("Statement - Black")
        put_title(deck, slide, 100, spec, f"{path}.title", size=32, anchor="m")
        support = text_of(require(spec, "support", f"{path}.support"), f"{path}.support", 140)
        fill_named(deck, slide, 101, support, f"{path}.support", 16, K["white"])
        arch(slide, 11.15, 5.2, 0.62, 1.35, K["lime"])
        arch(slide, 11.9, 4.55, 0.55, 2.0, K["sky"])
        return slide
    slide = deck.add("Title Only - White")
    put_title(deck, slide, 100, spec, f"{path}.title")
    stats = require(spec, "stats", f"{path}.stats")
    if not isinstance(stats, list) or not stats or len(stats) > 4:
        fail(f"{path}.stats", "1 to 4 stats")
    gap = 0.16
    width = (12.13 - gap * (len(stats) - 1)) / len(stats)
    for stat_index, stat in enumerate(stats):
        at = f"{path}.stats.{stat_index}"
        if not isinstance(stat, dict):
            fail(at, "expected an object")
        label = with_pace(text_of(require(stat, "label", f"{at}.label"), f"{at}.label", 24))
        value = key_of(require(stat, "key", f"{at}.key"), f"{at}.key")
        x = 0.6 + stat_index * (width + gap)
        card(slide, x, 2.0, width, 2.22, K["sand"])
        numdot(slide, x + 0.14, 2.14, stat_index + 1, K["forest"], K["white"], 0.36)
        say(slide, x + 0.14, 2.58, width - 0.28, 0.7, value, f"{at}.key", size=22 if len(value) > 12 else 26, color=K["forest"], font=HEAD)
        tall = len(label) > 18
        say(slide, x + 0.14, 3.28, width - 0.28, 0.48 if tall else 0.32, label, f"{at}.label", size=11 if tall else 12, color=K["ink"])
        if stat.get("note") is not None:
            note = text_of(stat["note"], f"{at}.note", 48)
            say(slide, x + 0.14, 3.78 if tall else 3.6, width - 0.28, 0.38, note, f"{at}.note", size=11, color=K["muted"])
    outcome = text_of(require(spec, "outcome", f"{path}.outcome"), f"{path}.outcome", 160)
    text_w = 10.1
    top = 4.22 + 0.3
    height = fitted_copy_height(outcome, text_w, 16)
    card(slide, 0.6, top, 12.13, height, K["sand"])
    tag(slide, 0.82, top + (height / 2) - 0.16, "Outcome", K["lime"], K["ink"])
    shape = say(slide, 2.35, top + 0.3, text_w, height - 0.6, outcome, f"{path}.outcome", size=16, color=K["ink"], anchor="t")
    shape.text_frame.margin_top = 0
    shape.text_frame.margin_bottom = 0
    return slide


def bars(slide, x, y, width, label, value, path, max_value):
    say(slide, x, y, width * 0.62, 0.28, label, path + ".label", size=12, color=K["ink"])
    say(slide, x + width * 0.62, y, width * 0.38, 0.28, value, path + ".value", size=12, color=K["forest"], font=HEAD, align="r")
    pill(slide, x, y + 0.32, width, 0.18, K["white"])
    ratio = 0 if max_value <= 0 else min(1, magnitude(value) / max_value)
    pill(slide, x, y + 0.32, max(0.18, width * ratio), 0.18, K["forest"])


def build_channel(deck, spec, index, photos):
    path = f"slides.{index}"
    kind = variant(spec, path, ("bars", "cards", "photo", "chart"))
    if kind == "chart":
        return build_delivery_layout(deck, spec, index, fixed=True)
    if kind == "bars":
        slide = deck.add("Title Only - Sand")
        put_title(deck, slide, 100, spec, f"{path}.title")
        spend = key_of(require(spec, "spend", f"{path}.spend"), f"{path}.spend")
        expected = key_of(require(spec, "expected", f"{path}.expected"), f"{path}.expected")
        spend_label = text_of(spec.get("spendLabel") or "Spend", f"{path}.spendLabel", 24)
        expected_label = text_of(spec.get("expectedLabel") or "Expected", f"{path}.expectedLabel", 24)
        cap = max(magnitude(spend), magnitude(expected), 1)
        bars(slide, 0.6, 1.95, 12.13, spend_label, spend, f"{path}.spend", cap)
        bars(slide, 0.6, 2.62, 12.13, expected_label, expected, f"{path}.expected", cap)
        delivery = require(spec, "delivery", f"{path}.delivery")
        if not isinstance(delivery, list) or not delivery or len(delivery) > 4:
            fail(f"{path}.delivery", "1 to 4 delivery figures")
        gap = 0.14
        card_w = (12.13 - gap * (len(delivery) - 1)) / len(delivery)
        for item_index, item in enumerate(delivery):
            at = f"{path}.delivery.{item_index}"
            if not isinstance(item, dict):
                fail(at, "expected an object")
            label = with_pace(text_of(require(item, "label", f"{at}.label"), f"{at}.label", 24))
            value = key_of(require(item, "key", f"{at}.key"), f"{at}.key")
            x = 0.6 + item_index * (card_w + gap)
            card(slide, x, 3.5, card_w, 1.28, K["white"])
            say(slide, x + 0.12, 3.58, card_w - 0.24, 0.55, value, f"{at}.key", size=14 if len(value) > 14 else 18, color=K["forest"], font=HEAD)
            say(slide, x + 0.12, 4.16, card_w - 0.24, 0.42, label, f"{at}.label", size=12, color=K["muted"])
        insight = text_of(require(spec, "insight", f"{path}.insight"), f"{path}.insight", 160)
        card(slide, 0.6, 4.98, 12.13, 1.52, K["white"])
        tag(slide, 0.82, 5.16, "Insight", K["sky"], K["ink"])
        say(slide, 2.35, 5.12, 10.1, 1.2, insight, f"{path}.insight", size=14, color=K["ink"], anchor="m")
        return slide
    if kind == "cards":
        slide = deck.add("Title Only - White")
        put_title(deck, slide, 100, spec, f"{path}.title")
        pointer = require(spec, "channel", f"{path}.channel")
        channel_index = channel_index_from_pointer(pointer, f"{path}.channel")
        key_of(pointer, f"{path}.channel")
        cards = require(spec, "cards", f"{path}.cards")
        if not isinstance(cards, list) or len(cards) != 4:
            fail(f"{path}.cards", "exactly 4 cards")
        drawn = []
        for item_index, item in enumerate(cards):
            at = f"{path}.cards.{item_index}"
            if not isinstance(item, dict):
                fail(at, "expected an object")
            label, field = FIXED_STATS[item_index]
            fill, ink = DOTS[item_index]
            drawn.append({
                "title": key_of(f"channels.{channel_index}.{field}", f"{path}.{field}"),
                "body": text_of(require(item, "body", f"{at}.body"), f"{at}.body", 120),
                "tag": label,
                "fill": fill,
                "ink": ink,
                "tag_fill": K["forest"],
                "tag_ink": K["white"],
                "title_path": f"{path}.{field}",
                "body_path": f"{at}.body",
                "title_color": K["forest"],
            })
        height = draw_fit_cards(slide, drawn, 1.95, 0.6, 12.73)
        insight_card(slide, path, spec, 1.95 + height)
        return slide
    slide = deck.add("Photo and Text")
    put_title(deck, slide, 100, spec, f"{path}.title", size=32)
    value = key_of(require(spec, "value", f"{path}.value"), f"{path}.value")
    body = text_of(require(spec, "body", f"{path}.body"), f"{path}.body", 200)
    fill_named(
        deck, slide, 101,
        [[Run(value, font=HEAD, color=K["forest"], size=32)], [body]],
        f"{path}.body", 16, K["ink"],
    )
    photo = photo_at(photos, spec.get("photo", 0), f"{path}.photo")
    deck.photo(slide, 102, photo, kind="arch")
    return slide


def build_not_live(deck, spec, index):
    path = f"slides.{index}"
    kind = variant(spec, path, ("cards", "list"))
    slide = deck.add("Title Only - White")
    put_title(deck, slide, 100, spec, f"{path}.title")
    cards = None
    from_list = False
    if kind == "list":
        rows = require(spec, "rows", f"{path}.rows")
        if not isinstance(rows, list) or not rows or len(rows) > 5:
            fail(f"{path}.rows", "1 to 5 rows")
        if len(rows) <= 2:
            cards = rows
            from_list = True
            kind = "cards"
    if kind == "cards":
        if cards is None:
            cards = require(spec, "cards", f"{path}.cards")
        if not isinstance(cards, list) or not cards or len(cards) > 4:
            fail(f"{path}.cards", "1 to 4 cards")
        gap = 0.18
        width = 6.4 if len(cards) == 1 else (12.13 - gap * (len(cards) - 1)) / len(cards)
        if len(cards) == 1:
            arch(slide, 8.15, 3.55, 1.05, 1.9, K["sky"])
            arch(slide, 9.4, 2.85, 1.1, 2.6, K["forest"])
            arch(slide, 10.7, 2.15, 1.15, 3.3, K["lime"])
        for item_index, item in enumerate(cards):
            at = f"{path}.cards.{item_index}"
            if not isinstance(item, dict):
                fail(at, "expected an object")
            label = text_of(require(item, "label", f"{at}.label"), f"{at}.label", 32)
            budget = key_of(require(item, "budget", f"{at}.budget"), f"{at}.budget")
            status = text_of(item["status"], f"{at}.status", 48) if item.get("status") is not None else None
            note = text_of(item["note"], f"{at}.note", 80) if item.get("note") is not None else None
            if from_list and note is None:
                fail(at, "missing note")
            if not from_list and status is None:
                fail(at, "missing status")
            x = 0.6 + item_index * (width + gap)
            card(slide, x, 2.15, width, 3.55, K["sand"])
            tag(slide, x + 0.2, 2.32, "Planned", K["lime"], K["ink"])
            say(slide, x + 0.2, 2.82, width - 0.4, 0.36, label, f"{at}.label", size=16, color=K["ink"], font=HEAD)
            say(slide, x + 0.2, 3.24, width - 0.4, 0.55, budget, f"{at}.budget", size=22 if len(budget) > 12 else 28, color=K["forest"], font=HEAD)
            y = 3.9
            if status is not None:
                say(slide, x + 0.2, y, width - 0.4, 0.55, status, f"{at}.status", size=14, color=K["body"])
                y += 0.6
            if note is not None:
                say(slide, x + 0.2, y, width - 0.4, 0.7, note, f"{at}.note", size=14, color=K["body"])
        return slide
    rows = require(spec, "rows", f"{path}.rows")
    if not isinstance(rows, list) or not rows or len(rows) > 5:
        fail(f"{path}.rows", "1 to 5 rows")
    for item_index, item in enumerate(rows):
        at = f"{path}.rows.{item_index}"
        if not isinstance(item, dict):
            fail(at, "expected an object")
        label = text_of(require(item, "label", f"{at}.label"), f"{at}.label", 32)
        budget = key_of(require(item, "budget", f"{at}.budget"), f"{at}.budget")
        note = text_of(require(item, "note", f"{at}.note"), f"{at}.note", 80)
        y = 2.15 + item_index * 0.85
        numdot(slide, 0.6, y, item_index + 1, K["forest"], K["white"], 0.42)
        say(slide, 1.25, y, 4.2, 0.42, label, f"{at}.label", size=16, color=K["ink"], font=HEAD, anchor="m")
        say(slide, 5.6, y, 3.2, 0.42, budget, f"{at}.budget", size=16, color=K["forest"], font=HEAD, anchor="m")
        say(slide, 8.9, y, 3.8, 0.42, note, f"{at}.note", size=14, color=K["body"], anchor="m")
        line(slide, 1.25, y + 0.62, 11.4, K["line"], 1.25)
    return slide


def build_kpi(deck, spec, index):
    path = f"slides.{index}"
    kind = variant(spec, path, ("four-cards", "table"))
    if kind == "table":
        slide = deck.add("Title Only - White")
        put_title(deck, slide, 100, spec, f"{path}.title")
        rows = require(spec, "rows", f"{path}.rows")
        if not isinstance(rows, list) or not rows:
            fail(f"{path}.rows", "at least one row")
        table_rows = [["KPI", "Target", "Actual"]]
        for item_index, item in enumerate(rows):
            at = f"{path}.rows.{item_index}"
            if not isinstance(item, dict):
                fail(at, "expected an object")
            label = key_of(require(item, "label", f"{at}.label"), f"{at}.label")
            table_rows.append([
                label,
                kpi_figure(label, key_of(require(item, "target", f"{at}.target"), f"{at}.target")),
                kpi_figure(label, key_of(require(item, "actual", f"{at}.actual"), f"{at}.actual")),
            ])
        grid = table(slide, 0.6, 2.15, 12.13, table_rows, widths=[3, 4, 4], size=13, head=True, first_bold=True, row_h=0.42)
        grid.name = f"{path}.table"
        line(slide, 0.6, 2.05, 2.2, K["forest"], 2.5)
        above = 2.15 + 0.42 * len(table_rows)
        if any(kpi_metric_is_conversion(item) for item in rows if isinstance(item, dict)):
            say(
                slide, 0.6, above + 0.06, 12.13, 0.22,
                CONVERSION_NOTE, f"{path}.conversion", size=9, color=K["muted"],
            )
            above += 0.28
        insight_card(slide, path, spec, above)
        return slide
    slide = deck.add("Four Cards")
    put_title(deck, slide, 100, spec, f"{path}.title")
    cards = require(spec, "cards", f"{path}.cards")
    if not isinstance(cards, list) or not cards or len(cards) > 4:
        fail(f"{path}.cards", "1 to 4 cards")
    slots = ((103, 104), (107, 108), (111, 112), (115, 116))
    for item_index, (head, body) in enumerate(slots):
        if item_index >= len(cards):
            deck.drop(slide, head, body)
            continue
        item = cards[item_index]
        at = f"{path}.cards.{item_index}"
        if not isinstance(item, dict):
            fail(at, "expected an object")
        label = key_of(require(item, "label", f"{at}.label"), f"{at}.label")
        actual = kpi_figure(label, key_of(require(item, "actual", f"{at}.actual"), f"{at}.actual"))
        target = kpi_figure(label, key_of(require(item, "target", f"{at}.target"), f"{at}.target"))
        fill_named(deck, slide, head, label, f"{at}.label", 14, K["ink"], HEAD)
        fill_named(
            deck, slide, body,
            [[Run(actual, font=HEAD, color=K["forest"], size=20)], [f"Target {target}"]],
            f"{at}.actual", 12, K["muted"],
        )
    return slide


def build_iao(deck, spec, index):
    path = f"slides.{index}"
    kind = variant(spec, path, ("rows", "journey"))
    if kind == "journey":
        slide = deck.add("Title Only - White")
        put_title(deck, slide, 100, spec, f"{path}.title")
        steps = require(spec, "steps", f"{path}.steps")
        if not isinstance(steps, list) or len(steps) != 3:
            fail(f"{path}.steps", "exactly 3 steps")
        colours = (K["sky"], K["forest"], K["lime"])
        tags = ("Insight", "Action", "Outcome")
        journey(slide, 0.9, 2.25, 11.5, colours, d=0.46, lc=K["ink"])
        for item_index, item in enumerate(steps):
            at = f"{path}.steps.{item_index}"
            if not isinstance(item, dict):
                fail(at, "expected an object")
            label = text_of(require(item, "label", f"{at}.label"), f"{at}.label", 24)
            copy = text_of(require(item, "text", f"{at}.text"), f"{at}.text", words=35)
            x = 0.6 + item_index * 4.15
            tag(slide, x, 3.0, tags[item_index], colours[item_index], K["ink"] if item_index != 1 else K["white"])
            say(slide, x, 3.5, 3.8, 0.4, label, f"{at}.label", size=16, color=K["ink"], font=HEAD)
            say(slide, x, 4.0, 3.8, 1.6, copy, f"{at}.text", size=14, color=K["body"])
        return slide
    slide = deck.add("Title Only - Black")
    put_title(deck, slide, 100, spec, f"{path}.title")
    rows = require(spec, "rows", f"{path}.rows")
    if not isinstance(rows, list) or not rows or len(rows) > 4:
        fail(f"{path}.rows", "1 to 4 rows")
    colours = {"insight": (K["sky"], K["ink"], "Insight"), "action": (K["forest"], K["white"], "Action"), "outcome": (K["lime"], K["ink"], "Outcome")}
    for item_index, item in enumerate(rows):
        at = f"{path}.rows.{item_index}"
        if not isinstance(item, dict):
            fail(at, "expected an object")
        tag_name = item.get("tag")
        if tag_name not in colours:
            fail(f"{at}.tag", "insight, action or outcome")
        copy = text_of(require(item, "text", f"{at}.text"), f"{at}.text", words=35)
        fill, ink, label = colours[tag_name]
        y = 2.05 + item_index * 1.1
        tag(slide, 0.6, y, label, fill, ink)
        say(slide, 2.35, y, 10.2, 0.85, copy, f"{at}.text", size=15, color=K["white"], anchor="m")
    return slide


def build_appendix(deck, spec, index):
    path = f"slides.{index}"
    if spec.get("variant") not in (None, "table"):
        fail(path, "variant must be table")
    slide = deck.add("Title Only - White")
    put_title(deck, slide, 100, spec, f"{path}.title")
    columns = require(spec, "columns", f"{path}.columns")
    rows = require(spec, "rows", f"{path}.rows")
    if not isinstance(columns, list) or not columns:
        fail(f"{path}.columns", "at least one column")
    if not isinstance(rows, list) or not rows:
        fail(f"{path}.rows", "at least one row")
    header = []
    for col_index, column in enumerate(columns):
        header.append(text_of(column, f"{path}.columns.{col_index}", 18))
    table_rows = [header]
    width = len(header)
    for row_index, row in enumerate(rows):
        if not isinstance(row, list) or len(row) != width:
            fail(f"{path}.rows.{row_index}", f"{width} keys")
        table_rows.append([key_of(cell, f"{path}.rows.{row_index}.{cell_index}") for cell_index, cell in enumerate(row)])
    grid = table(slide, 0.6, 2.15, 12.13, table_rows, size=12, head=True, first_bold=True, row_h=0.42)
    grid.name = f"{path}.table"
    line(slide, 0.6, 2.05, 2.2, K["forest"], 2.5)
    insight_card(slide, path, spec, 2.15 + 0.42 * len(table_rows))
    return slide


def clear_slide_number(slide):
    """Drop the example slide's cached page number. The field stays."""
    for node in slide._element.iter(qn("a:fld")):
        if node.get("type") != "slidenum":
            continue
        for text in node.iter(qn("a:t")):
            text.text = ""


def build_breaker(deck, spec, index, photos):
    """Black breaker cloned from template slide 5.

    The clone keeps the dark ground and the bottom-right bar motif
    (sky, forest light, lime). The title is filled, not drawn.
    """
    path = f"slides.{index}"
    variant(spec, path, ("black",))
    slide = deck.clone(5)
    clear_slide_number(slide)
    visible, runs = parse_title(require(spec, "title", f"{path}.title"), f"{path}.title")
    size = 32 if len(visible) > 36 else 36
    shape = deck.fill(
        slide, 102, runs,
        base=dict(size=size, color=fg(slide), font=HEAD),
        align="l",
    )
    no_autofit(shape.text_frame)
    shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    shape.name = f"{path}.title"
    deck.drop(slide, 101)
    deck.photo(slide, 100, photos[0], kind="arch")
    return slide


def bar_motif(slide, x=10.10, bottom=6.70):
    """Three arches, sky then forest then lime, feet on one baseline.

    On a dark slide the middle arch is forest light.
    """
    sky_h, forest_h, lime_h = 1.76, 2.48, 3.36
    middle = K["fl"] if slide._am_dark else K["forest"]
    arch(slide, x, bottom - sky_h, 0.72, sky_h, K["sky"])
    arch(slide, x + 0.84, bottom - forest_h, 0.72, forest_h, middle)
    arch(slide, x + 1.68, bottom - lime_h, 0.72, lime_h, K["lime"])


def kpi_metric_is_conversion(row):
    if not isinstance(row, dict):
        return False
    if row.get("metric") == "conversion_rate":
        return True
    if not isinstance(DATA, dict):
        return False
    kpis = DATA.get("kpis")
    if not isinstance(kpis, list):
        return False
    for field in ("label", "actual", "target"):
        key = row.get(field)
        if not isinstance(key, str):
            continue
        match = re.fullmatch(r"kpis\.(\d+)\..+", key)
        if not match:
            continue
        index = int(match.group(1))
        if index < len(kpis) and isinstance(kpis[index], dict) and kpis[index].get("metric") == "conversion_rate":
            return True
    return False


def build_chart(deck, spec, index):
    variant(spec, f"slides.{index}", ("delivery",))
    return build_delivery_layout(deck, spec, index, fixed=False)


def build_recommendations(deck, spec, index):
    path = f"slides.{index}"
    variant(spec, path, ("cards",))
    slide = deck.add("Title Only - White")
    put_title(deck, slide, 100, spec, f"{path}.title")
    cards = require(spec, "cards", f"{path}.cards")
    if not isinstance(cards, list) or len(cards) not in (2, 3):
        fail(f"{path}.cards", "2 or 3 recommendations")
    outcome = text_of(require(spec, "outcome", f"{path}.outcome"), f"{path}.outcome", 160)
    drawn = []
    for item_index, item in enumerate(cards):
        at = f"{path}.cards.{item_index}"
        if not isinstance(item, dict):
            fail(at, "expected an object")
        fill, ink = DOTS[item_index]
        drawn.append({
            "title": text_of(require(item, "heading", f"{at}.heading"), f"{at}.heading", 40),
            "body": text_of(require(item, "body", f"{at}.body"), f"{at}.body", 220),
            "tag": "Action",
            "fill": fill,
            "ink": ink,
            "tag_fill": K["forest"],
            "tag_ink": K["white"],
            "title_path": f"{at}.heading",
            "body_path": f"{at}.body",
        })
    right = 9.85 if len(cards) == 2 else 12.73
    height = draw_fit_cards(slide, drawn, 1.95, 0.6, right)
    cards_bottom = 1.95 + height
    if len(cards) == 2:
        bar_motif(slide, bottom=cards_bottom)
    outcome_top = cards_bottom + 0.35
    tag(slide, 0.6, outcome_top, "Outcome", K["lime"], K["ink"])
    say(slide, 2.2, outcome_top, 10.5, 0.4, outcome, f"{path}.outcome", size=15, color=K["ink"], anchor="t")
    return slide


def build_close(deck, spec, index, photos):
    path = f"slides.{index}"
    kind = variant(spec, path, ("thank-you", "cta"))
    if kind == "thank-you":
        slide = deck.clone(21)
        visible, runs = parse_title(require(spec, "title", f"{path}.title"), f"{path}.title")
        size = 32 if len(visible) > 28 else 36
        shape = deck.fill(slide, 101, runs, base=dict(size=size, color=K["white"], font=HEAD), align="l")
        no_autofit(shape.text_frame)
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        shape.name = f"{path}.title"
        deck.photo(slide, 103, photos[0], kind="arch")
        return slide
    slide = deck.add("Title Only - Black")
    put_title(deck, slide, 100, spec, f"{path}.title")
    support = text_of(require(spec, "support", f"{path}.support"), f"{path}.support", 140)
    button = text_of(require(spec, "button", f"{path}.button"), f"{path}.button", 24)
    say(slide, 0.6, 2.3, 7.6, 1.4, support, f"{path}.support", size=18, color=K["white"])
    btn(slide, 0.6, 4.3, button, primary=True)
    trio(slide, 8.4, 6.55, 0.95, dark=True)
    return slide


BUILDERS = {
    "cover": build_cover,
    "breaker": build_breaker,
    "summary": build_summary,
    "chart": build_chart,
    "channel": build_channel,
    "not-live": build_not_live,
    "kpi": build_kpi,
    "iao": build_iao,
    "recommendations": build_recommendations,
    "appendix": build_appendix,
    "close": build_close,
}


def not_live_lookup(data):
    """Index and group names for channels the payload marks not live."""
    indexes = {}
    groups = {}
    listed = {}
    if not isinstance(data, dict):
        return indexes, groups, listed
    channels = data.get("channels")
    if isinstance(channels, list):
        for index, channel in enumerate(channels):
            if not isinstance(channel, dict) or channel.get("notLive") is not True:
                continue
            label = channel.get("label") if isinstance(channel.get("label"), str) else channel.get("group")
            name = label if isinstance(label, str) else f"channels.{index}"
            indexes[str(index)] = name
            group = channel.get("group")
            if isinstance(group, str):
                groups[group] = name
    rows = data.get("notLive")
    if isinstance(rows, list):
        for index, channel in enumerate(rows):
            if not isinstance(channel, dict):
                continue
            label = channel.get("label") if isinstance(channel.get("label"), str) else channel.get("group")
            name = label if isinstance(label, str) else f"notLive.{index}"
            listed[index] = name
            group = channel.get("group")
            if isinstance(group, str) and group not in groups:
                groups[group] = name
    return indexes, groups, listed


def not_live_name(key, indexes, groups, listed):
    if key.startswith("notLive."):
        parts = key.split(".")
        if len(parts) >= 2 and parts[1].isdigit():
            return listed.get(int(parts[1]), "notLive channel")
        return "notLive channel"
    match = re.fullmatch(r"channels\.([A-Za-z0-9_]+)\..+", key)
    if not match:
        return None
    part = match.group(1)
    if part in indexes:
        return indexes[part]
    if part in groups:
        return groups[part]
    return None


def channel_token(value):
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def channel_field_name(value, indexes, groups, listed):
    """The not-live channel named by a slide's own channel field."""
    if not isinstance(value, str) or not value.strip():
        return None
    if value in groups:
        return groups[value]
    if value in indexes:
        return indexes[value]
    wanted = channel_token(value)
    if not wanted:
        return None
    for name in list(groups.values()) + list(indexes.values()) + list(listed.values()):
        folded = channel_token(name)
        if folded == wanted or folded.endswith(wanted):
            return name
    for group, name in groups.items():
        if channel_token(group) == wanted:
            return name
    return None


def primary_channel_keys(spec):
    """Value keys that say which channel this slide is about."""
    found = []
    kind = spec.get("type")
    if kind == "channel":
        for field in ("spend", "expected", "value", "channel"):
            node = spec.get(field)
            if isinstance(node, str) and node.startswith("channels."):
                found.append(node)
        for group in ("delivery", "cards"):
            rows = spec.get(group)
            if not isinstance(rows, list):
                continue
            for item in rows:
                if isinstance(item, dict) and isinstance(item.get("key"), str):
                    found.append(item["key"])
    elif kind == "not-live":
        for group in ("cards", "rows"):
            rows = spec.get(group)
            if not isinstance(rows, list):
                continue
            for item in rows:
                if isinstance(item, dict) and isinstance(item.get("budget"), str):
                    found.append(item["budget"])
    return found


def not_live_slide_error(number, name):
    return (
        f"slide {number} (channel): {name} is not live. "
        f"Remove this slide and mention {name} in the summary Outcome instead."
    )


def validate_content(content, data=None):
    """Schema checks. Resolved string length is checked at build.

    A not-live channel has no slide of its own. Summary, KPI and rates
    slides may still name it. The slide's own channel is its channel field
    or, on a channel or not-live slide, its primary value keys.
    """
    errors = []

    def bad(path, message):
        errors.append(f"{path}: {message}")

    def prose(node, path, limit=None, words=None):
        if node is None:
            bad(path, "missing")
            return
        if isinstance(node, str) and KEY_RE.fullmatch(node):
            return
        if isinstance(node, str):
            if digits(node):
                bad(path, "digits in copy; use a value key")
            elif limit is not None and len(node) > limit:
                bad(path, f"{len(node)} characters, limit {limit}")
            if words is not None and len(re.findall(r"\S+", node)) > words:
                bad(path, f"over {words} words")
            return
        if isinstance(node, list):
            length = 0
            count = 0
            for index, part in enumerate(node):
                at = f"{path}.{index}"
                if isinstance(part, str):
                    if digits(part):
                        bad(at, "digits in copy; use a value key")
                    length += len(part)
                    count += len(re.findall(r"\S+", part))
                elif isinstance(part, dict) and isinstance(part.get("k"), str) and len(part) == 1:
                    if not KEY_RE.fullmatch(part["k"]):
                        bad(f"{at}.k", "expected a value key")
                    count += 1
                else:
                    bad(at, 'expected copy or {"k": path}')
            if limit is not None and length > limit:
                bad(path, f"{length} characters, limit {limit}")
            if words is not None and count > words:
                bad(path, f"over {words} words")
            return
        bad(path, "expected copy or a value key")

    def key(node, path):
        if not isinstance(node, str) or not KEY_RE.fullmatch(node):
            bad(path, "expected a value key")

    def title(spec, path):
        if "title" not in spec:
            bad(f"{path}.title", "missing title")
            return
        try:
            parse_title(spec["title"], f"{path}.title")
        except BuildError as error:
            errors.append(str(error))

    def variant_name(spec, path, allowed):
        name = spec.get("variant")
        if name not in allowed:
            bad(path, f"variant must be {' or '.join(allowed)}")
            return None
        return name

    def items(node, path, low, high, noun):
        if not isinstance(node, list) or len(node) < low or (high is not None and len(node) > high):
            span = str(low) if low == high else (f"{low} to {high}" if high is not None else f"at least {low}")
            bad(path, f"{span} {noun}")
            return []
        return node

    if not isinstance(content, dict) or not isinstance(content.get("slides"), list) or not content["slides"]:
        return ["slides: expected a slide list"]

    indexes, groups, listed = not_live_lookup(data)

    for index, spec in enumerate(content["slides"]):
        path = f"slides.{index}"
        if not isinstance(spec, dict):
            bad(path, "expected an object")
            continue
        kind = spec.get("type")
        if kind not in BUILDERS:
            bad(path, "unknown slide type")
            continue
        title(spec, path)
        if kind == "cover":
            name = variant_name(spec, path, ("sand", "black"))
            if name is not None:
                key(spec.get("subtitle"), f"{path}.subtitle")
        elif kind == "summary":
            name = variant_name(spec, path, ("stat-row", "statement"))
            if name == "statement":
                prose(spec.get("support"), f"{path}.support", 140)
            elif name == "stat-row":
                for item_index, item in enumerate(items(spec.get("stats"), f"{path}.stats", 1, 4, "stats")):
                    at = f"{path}.stats.{item_index}"
                    if not isinstance(item, dict):
                        bad(at, "expected an object")
                        continue
                    prose(item.get("label"), f"{at}.label", 24)
                    key(item.get("key"), f"{at}.key")
                    if item.get("note") is not None:
                        prose(item.get("note"), f"{at}.note", 48)
                prose(spec.get("outcome"), f"{path}.outcome", 160)
        elif kind == "chart":
            name = variant_name(spec, path, ("delivery",))
            if name == "delivery":
                chart_key = spec.get("chart")
                key(chart_key, f"{path}.chart")
                if not isinstance(data, dict) or as_chart(data, chart_key) is None:
                    bad(f"{path}.chart", "chart must resolve to an object with points. Point chart at charts.0.")
                stats = spec.get("stats")
                if not isinstance(stats, list) or len(stats) > 4:
                    bad(f"{path}.stats", "up to 4 stats")
                else:
                    for item_index, item in enumerate(stats):
                        at = f"{path}.stats.{item_index}"
                        if not isinstance(item, dict):
                            bad(at, "expected an object")
                            continue
                        prose(item.get("label"), f"{at}.label", 24)
                        key(item.get("key"), f"{at}.key")
                prose(spec.get("insight"), f"{path}.insight", 160)
        elif kind == "channel":
            name = variant_name(spec, path, ("bars", "cards", "photo", "chart"))
            if name == "bars":
                key(spec.get("spend"), f"{path}.spend")
                key(spec.get("expected"), f"{path}.expected")
                if spec.get("spendLabel") is not None:
                    prose(spec.get("spendLabel"), f"{path}.spendLabel", 24)
                if spec.get("expectedLabel") is not None:
                    prose(spec.get("expectedLabel"), f"{path}.expectedLabel", 24)
                for item_index, item in enumerate(items(spec.get("delivery"), f"{path}.delivery", 1, 4, "delivery figures")):
                    at = f"{path}.delivery.{item_index}"
                    if not isinstance(item, dict):
                        bad(at, "expected an object")
                        continue
                    prose(item.get("label"), f"{at}.label", 24)
                    key(item.get("key"), f"{at}.key")
                prose(spec.get("insight"), f"{path}.insight", 160)
            elif name == "cards":
                pointer = spec.get("channel")
                if not isinstance(pointer, str) or CHANNEL_KEY_RE.fullmatch(pointer) is None:
                    bad(f"{path}.channel", "expected a channels.<i> key, such as channels.0.label")
                else:
                    key(pointer, f"{path}.channel")
                cards = items(spec.get("cards"), f"{path}.cards", 4, 4, "cards")
                for item_index, item in enumerate(cards):
                    at = f"{path}.cards.{item_index}"
                    if not isinstance(item, dict):
                        bad(at, "expected an object")
                        continue
                    prose(item.get("body"), f"{at}.body", 120)
                prose(spec.get("insight"), f"{path}.insight", 160)
            elif name == "chart":
                key(spec.get("chart"), f"{path}.chart")
                problem = channel_chart_problem(spec, data)
                if problem:
                    where = path if problem == "use variant cards" else f"{path}.chart"
                    bad(where, problem)
                prose(spec.get("insight"), f"{path}.insight", 160)
            elif name == "photo":
                key(spec.get("value"), f"{path}.value")
                prose(spec.get("body"), f"{path}.body", 200)
                photo = spec.get("photo", 0)
                if not isinstance(photo, int) or isinstance(photo, bool) or photo not in (0, 1, 2):
                    bad(f"{path}.photo", "photo index out of range")
        elif kind == "not-live":
            name = variant_name(spec, path, ("cards", "list"))
            if name == "cards":
                for item_index, item in enumerate(items(spec.get("cards"), f"{path}.cards", 1, 4, "cards")):
                    at = f"{path}.cards.{item_index}"
                    if not isinstance(item, dict):
                        bad(at, "expected an object")
                        continue
                    prose(item.get("label"), f"{at}.label", 32)
                    key(item.get("budget"), f"{at}.budget")
                    prose(item.get("status"), f"{at}.status", 48)
                    if item.get("note") is not None:
                        prose(item.get("note"), f"{at}.note", 80)
            elif name == "list":
                rows = items(spec.get("rows"), f"{path}.rows", 1, 5, "rows")
                if rows and len(rows) <= 2:
                    bad(f"{path}.rows", "list is not allowed for 1 or 2 channels; use cards")
                for item_index, item in enumerate(rows):
                    at = f"{path}.rows.{item_index}"
                    if not isinstance(item, dict):
                        bad(at, "expected an object")
                        continue
                    prose(item.get("label"), f"{at}.label", 32)
                    key(item.get("budget"), f"{at}.budget")
                    prose(item.get("note"), f"{at}.note", 80)
        elif kind == "kpi":
            name = variant_name(spec, path, ("four-cards", "table"))
            group = "cards" if name == "four-cards" else "rows"
            if name is not None:
                entries = items(spec.get(group), f"{path}.{group}", 1, 4 if name == "four-cards" else None, group)
                for item_index, item in enumerate(entries):
                    at = f"{path}.{group}.{item_index}"
                    if not isinstance(item, dict):
                        bad(at, "expected an object")
                        continue
                    key(item.get("label"), f"{at}.label")
                    key(item.get("actual"), f"{at}.actual")
                    key(item.get("target"), f"{at}.target")
                if name == "table":
                    prose(spec.get("insight"), f"{path}.insight", 160)
        elif kind == "iao":
            name = variant_name(spec, path, ("rows", "journey"))
            if name == "journey":
                for item_index, item in enumerate(items(spec.get("steps"), f"{path}.steps", 3, 3, "steps")):
                    at = f"{path}.steps.{item_index}"
                    if not isinstance(item, dict):
                        bad(at, "expected an object")
                        continue
                    prose(item.get("label"), f"{at}.label", 24)
                    prose(item.get("text"), f"{at}.text", words=35)
            elif name == "rows":
                for item_index, item in enumerate(items(spec.get("rows"), f"{path}.rows", 1, 4, "rows")):
                    at = f"{path}.rows.{item_index}"
                    if not isinstance(item, dict):
                        bad(at, "expected an object")
                        continue
                    if item.get("tag") not in ("insight", "action", "outcome"):
                        bad(f"{at}.tag", "insight, action or outcome")
                    prose(item.get("text"), f"{at}.text", words=35)
        elif kind == "appendix":
            if spec.get("variant") not in (None, "table"):
                bad(path, "variant must be table")
            columns = spec.get("columns")
            rows = spec.get("rows")
            if not isinstance(columns, list) or not columns:
                bad(f"{path}.columns", "at least one column")
                columns = []
            else:
                for col_index, column in enumerate(columns):
                    prose(column, f"{path}.columns.{col_index}", 18)
            if not isinstance(rows, list) or not rows:
                bad(f"{path}.rows", "at least one row")
            else:
                for row_index, row in enumerate(rows):
                    if not isinstance(row, list) or len(row) != len(columns):
                        bad(f"{path}.rows.{row_index}", f"{len(columns)} keys")
                        continue
                    for cell_index, cell in enumerate(row):
                        key(cell, f"{path}.rows.{row_index}.{cell_index}")
            prose(spec.get("insight"), f"{path}.insight", 160)
        elif kind == "breaker":
            variant_name(spec, path, ("black",))
        elif kind == "recommendations":
            name = variant_name(spec, path, ("cards",))
            if name == "cards":
                entries = items(spec.get("cards"), f"{path}.cards", 2, 3, "recommendations")
                for item_index, item in enumerate(entries):
                    at = f"{path}.cards.{item_index}"
                    if not isinstance(item, dict):
                        bad(at, "expected an object")
                        continue
                    prose(item.get("heading"), f"{at}.heading", 40)
                    prose(item.get("body"), f"{at}.body", 220)
                prose(spec.get("outcome"), f"{path}.outcome", 160)
        elif kind == "close":
            name = variant_name(spec, path, ("thank-you", "cta"))
            if name == "cta":
                prose(spec.get("support"), f"{path}.support", 140)
                prose(spec.get("button"), f"{path}.button", 24)
        own = []
        named = channel_field_name(spec.get("channel"), indexes, groups, listed)
        if named:
            own.append(named)
        if kind in ("channel", "not-live"):
            for channel_key in primary_channel_keys(spec):
                named = not_live_name(channel_key, indexes, groups, listed)
                if named and named not in own:
                    own.append(named)
        for name in own:
            errors.append(not_live_slide_error(index + 1, name))
    return errors


def build(content, data, template, out, photos):
    global DATA
    DATA = data
    if not isinstance(content, dict) or not isinstance(content.get("slides"), list) or not content["slides"]:
        fail("slides", "expected a slide list")
    if len(photos) != 3:
        fail("photos", "expected 3 photos")
    for photo in photos:
        if not os.path.isfile(photo):
            fail("photos", f"missing photo {photo}")
    deck = Deck(template)
    for index, spec in enumerate(content["slides"]):
        path = f"slides.{index}"
        if not isinstance(spec, dict):
            fail(path, "expected an object")
        kind = spec.get("type")
        if kind not in BUILDERS:
            fail(path, "unknown slide type")
        if kind in ("cover", "breaker", "channel", "close"):
            slide = BUILDERS[kind](deck, spec, index, photos)
        else:
            slide = BUILDERS[kind](deck, spec, index)
        deck.notes(slide, f"CHECK slides.{index} {kind} {spec.get('variant', 'table')}")
    deck.save(out)
    return out


def main():
    if len(sys.argv) != 8:
        sys.stderr.write("usage: build_report.py content.json report-data.json template.potx out.pptx photo1 photo2 photo3\n")
        return 2
    content_path, data_path, template, out, *photos = sys.argv[1:]
    with open(content_path, encoding="utf-8") as handle:
        content = json.load(handle)
    with open(data_path, encoding="utf-8") as handle:
        data = json.load(handle)
    try:
        built = build(content, data, template, out, photos)
    except BuildError as error:
        sys.stderr.write(f"error: {error}\n")
        return 1
    import check_deck
    issues = check_deck.check(built, data_path)
    sys.stdout.write(check_deck.render(issues) + "\n")
    return 0 if not issues else 1


if __name__ == "__main__":
    sys.exit(main())
