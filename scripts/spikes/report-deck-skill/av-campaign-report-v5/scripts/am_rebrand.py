"""Assembled Media rebrand helpers (python-pptx) for building decks on the v5 template.
Usage: from am_rebrand import *; d = Deck(template_path); s = d.add('Title Only - White'); ...; d.save(out)
"""
import copy, os, re, zipfile, io
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image

K = dict(ink='0F1D13', white='FFFFFF', sand='EFE9DF', lime='B5D337', sky='49C7EB', forest='246646', fl='4E8F6A',
         panel='1A2620', ctx='CFC8BA', line='E4DED2', muted='6B6A5E', mutedB='AEB9B0', soft='C9D3C8', body='3B4A3F', grid='2A3A2F')
HEAD, BODY, SERIF = 'Plus Jakarta Sans ExtraBold', 'Plus Jakarta Sans', 'Instrument Serif'
DARK_LAYOUTS = ('Black', 'Statement - Black', 'Thank You')

def rgb(h): return RGBColor.from_string(h)

# ---------- runs: a str, or a list of str / I('serif italic') / B('bold') / C('text', 'hex') ----------
class Run:
    def __init__(self, t, **k): self.t, self.k = t, k
def I(t, **k): return Run(t, serif=True, **k)
def B(t, **k): return Run(t, bold=True, **k)
def C(t, color, **k): return Run(t, color=color, **k)

def _style_run(r, run, base):
    k = dict(base); k.update(run.k if isinstance(run, Run) else {})
    f = r.font
    if k.get('serif'): f.name = SERIF; f.italic = True; f.bold = False
    elif k.get('font'): f.name = k['font']
    if k.get('bold') is not None and not k.get('serif'): f.bold = k['bold']
    if k.get('size'): f.size = Pt(k['size'])
    if k.get('color'): f.color.rgb = rgb(k['color'])
    if k.get('spc') is not None:
        rPr = r._r.find(qn('a:rPr'))
        if rPr is not None: rPr.set('spc', str(int(k['spc'])))

def _run_text(run):
    return run.t if isinstance(run, Run) else run

def _copy_run(run, text, **extra):
    if isinstance(run, Run):
        k = dict(run.k); k.update(extra); return Run(text, **k)
    if extra: return Run(text, **extra)
    return text

def seal_stops(specs, base):
    """The closing stop belongs on the last run, with no space before it.

    Plus Jakarta's full stop sits wide of the letter. On a sans headline, pull that pair in.
    A serif stop is already tight, so it stays on its run untouched.
    """
    specs = [spec for spec in specs if _run_text(spec) != ""]
    if len(specs) >= 2 and isinstance(_run_text(specs[-1]), str) and re.fullmatch(r"[ \t]*[.!?]", _run_text(specs[-1])):
        punct = _run_text(specs[-1]).strip()
        specs[-2] = _copy_run(specs[-2], _run_text(specs[-2]).rstrip() + punct)
        specs.pop()
    sealed = []
    for spec in specs:
        text = _run_text(spec)
        if isinstance(text, str):
            closed = re.sub(r"[ \t]+([.!?])$", r"\1", text)
            if closed != text: spec = _copy_run(spec, closed)
        sealed.extend(_pull_sans_stop(spec, base))
    return [spec for spec in sealed if _run_text(spec) != ""]

def _pull_sans_stop(run, base):
    text = _run_text(run)
    k = run.k if isinstance(run, Run) else {}
    if k.get('serif') or not isinstance(text, str): return [run]
    font = k.get('font') or (base or {}).get('font')
    if font != HEAD: return [run]
    match = re.match(r"^(.*?)([.!?])$", text)
    if not match or not match.group(1) or not match.group(1)[-1].isalpha(): return [run]
    stem, punct = match.group(1), match.group(2)
    size = k.get('size') or (base or {}).get('size') or 16
    # About 0.12em: the ExtraBold stop's left sidebearing, in hundredths of a point.
    pull = -int(round(float(size) * 12))
    head, tail = stem[:-1], stem[-1] + punct
    tail_run = _copy_run(run, tail, spc=pull)
    if head: return [_copy_run(run, head), tail_run]
    return [tail_run]

def _runs(x): return [x] if isinstance(x, (str, Run)) else list(x)

def set_paras(tf, paras, base=None, bullet=False, space_after=None, align=None, line=None):
    """paras: a single run-spec, or list of paragraphs when bullet/list given as list of run-specs."""
    base = base or {}
    # str/Run -> one paragraph. List with nested lists -> paragraphs. List of plain strings (or bullet=True) -> one paragraph per item.
    # Any other list (strings mixed with I()/B() runs) -> one paragraph of runs.
    if isinstance(paras, (str, Run)): paras = [paras]
    elif any(isinstance(e, (list, tuple)) for e in paras): pass
    elif bullet or (len(paras) > 1 and all(isinstance(e, str) for e in paras)): pass
    else: paras = [paras]
    # clear
    for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    tmpl = None  # keep the first existing run's formatting (cloned slides carry their styling on the runs)
    if p0.runs and p0.runs[0]._r.find(qn('a:rPr')) is not None: tmpl = copy.deepcopy(p0.runs[0]._r.find(qn('a:rPr')))
    for r in list(p0.runs): r._r.getparent().remove(r._r)
    for i, spec in enumerate(paras):
        p = p0 if i == 0 else tf.add_paragraph()
        if align: p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if space_after is not None: p.space_after = Pt(space_after)
        if line: p.line_spacing = line
        if bullet:
            pPr = p._p.get_or_add_pPr(); pPr.set('marL', str(Emu(Inches(0.22)))); pPr.set('indent', str(-Emu(Inches(0.22))))
            for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum'):
                for e in pPr.findall(qn(tag)): pPr.remove(e)
            bu = etree.SubElement(pPr, qn('a:buChar')); bu.set('char', '●')
        for run in seal_stops(_runs(spec), base):
            r = p.add_run(); r.text = run.t if isinstance(run, Run) else run
            if tmpl is not None:
                old_rpr = r._r.find(qn('a:rPr'))
                if old_rpr is not None: r._r.remove(old_rpr)
                r._r.insert(0, copy.deepcopy(tmpl))
            _style_run(r, run, base)

class Deck:
    """Open the brand template, build new slides (from layouts or by cloning the template's example slides), then save.
    The template's own example slides are removed on save."""
    def __init__(self, template):
        data = open(template, 'rb').read()
        if template.endswith('.potx'):  # python-pptx needs the presentation content type
            zi = zipfile.ZipFile(io.BytesIO(data)); buf = io.BytesIO(); zo = zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED)
            for it in zi.infolist():
                b = zi.read(it)
                if it.filename == '[Content_Types].xml': b = b.replace(b'template.main+xml', b'presentation.main+xml')
                zo.writestr(it, b)
            zo.close(); buf.seek(0); data = buf.read()
        self.prs = Presentation(io.BytesIO(data))
        self.examples = list(self.prs.slides)            # template example slides (1-based index = position in template)
        self._ex_ids = [sid for sid in self.prs.slides._sldIdLst]
        self.layouts = {l.name: l for l in self.prs.slide_layouts}
        self.manual = []; self.count = 0

    def _register(self, s, dark):
        self.count += 1; s._am_no = self.count; s._am_dark = dark; return s

    def add(self, layout):
        """New slide from a layout. Copies the layout's slide-number placeholder so numbering shows."""
        lay = self.layouts[layout]; s = self.prs.slides.add_slide(lay)
        for sp in lay.placeholders:
            if sp.placeholder_format.type is not None and 'SLIDE_NUMBER' in str(sp.placeholder_format.type):
                s.shapes._spTree.append(copy.deepcopy(sp._element))
        return self._register(s, any(x in layout for x in DARK_LAYOUTS))

    def clone(self, n, keep_rings=False):
        """Copy example slide n (1-based, as numbered in the template) with all its shapes and images.
        Focus rings are removed by default: they were placed for the example photo, not yours. Add your own on the face or product."""
        src = self.examples[n - 1]; s = self.prs.slides.add_slide(src.slide_layout)
        for sh in list(s.shapes): sh._element.getparent().remove(sh._element)
        rid_map = {}
        for rid, rel in src.part.rels.items():
            if rel.is_external or 'slideLayout' in rel.reltype or 'notesSlide' in rel.reltype: continue
            rid_map[rid] = s.part.relate_to(rel._target, rel.reltype)
        spTree = s.shapes._spTree
        for el in src.shapes._spTree:
            if el.tag.endswith(('}sp', '}pic', '}grpSp', '}graphicFrame', '}cxnSp')):
                nm = el.find('.//' + qn('p:cNvPr'))
                if not keep_rings and nm is not None and nm.get('name') == 'Focus ring': continue
                e = copy.deepcopy(el)
                for node in e.iter():
                    for att in (qn('r:embed'), qn('r:link'), qn('r:id')):
                        v = node.get(att)
                        if v in rid_map: node.set(att, rid_map[v])
                spTree.append(e)
        bg = src._element.cSld.find(qn('p:bg'))
        if bg is not None: s._element.cSld.insert(0, copy.deepcopy(bg))
        return self._register(s, any(x in src.slide_layout.name for x in DARK_LAYOUTS))

    def ph(self, s, idx):
        for p in s.placeholders:
            if p.placeholder_format.idx == idx: return p
        raise KeyError(idx)

    def fill(self, s, idx, content, **kw):
        p = self.ph(s, idx); set_paras(p.text_frame, content, **kw); return p

    def drop(self, s, *idxs):
        """Remove unused placeholders so nothing empty is left on the slide."""
        for i in idxs:
            try: p = self.ph(s, i); p._element.getparent().remove(p._element)
            except KeyError: pass

    def photo(self, s, idx, path, kind='arch', focus=(0.5, 0.5)):
        """Swap a picture placeholder for a cropped image in the same spot and brand shape."""
        p = self.ph(s, idx); x, y, w, h = [v / 914400 for v in (p.left, p.top, p.width, p.height)]
        parent = p._element.getparent(); pos = list(parent).index(p._element); parent.remove(p._element)
        pic = image(s, path, x, y, w, h, kind, focus)
        parent.remove(pic._element); parent.insert(pos, pic._element)  # keep the placeholder's stacking order
        return pic

    def photo_slot(self, s, idx, what, kind='arch'):
        """Leave a picture placeholder empty for the team, flagged."""
        p = self.ph(s, idx); x, y, w, h = [v / 914400 for v in (p.left, p.top, p.width, p.height)]
        p._element.getparent().remove(p._element)
        return slot(self, s, x, y, w, h, kind, what)

    def flag(self, s, item): self.manual.append((s._am_no, item))

    def notes(self, s, text):
        ns = s.notes_slide
        if ns.notes_text_frame is None:  # template notes master has no body placeholder: add one
            from pptx.oxml import parse_xml
            sp = parse_xml('<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                '<p:nvSpPr><p:cNvPr id="90" name="Notes Placeholder"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
                '<p:spPr><a:xfrm><a:off x="685800" y="4400550"/><a:ext cx="5486400" cy="3600450"/></a:xfrm></p:spPr>'
                '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="en-AU"/></a:p></p:txBody></p:sp>')
            ns.shapes._spTree.append(sp)
        tf = ns.notes_text_frame; tf.text = (tf.text + '\n' if tf.text else '') + text

    def save(self, out):
        lst = self.prs.slides._sldIdLst
        for sid in self._ex_ids:
            self.prs.part.drop_rel(sid.get(qn('r:id'))); lst.remove(sid)
        slides = list(self.prs.slides)
        for no, item in self.manual: self.notes(slides[no - 1], 'MANUAL: ' + item)
        buf = io.BytesIO(); self.prs.save(buf); buf.seek(0)
        zi = zipfile.ZipFile(buf); names = zi.namelist(); zo = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
        for n in names:
            b = zi.read(n)
            if n == '[Content_Types].xml':  # declare any media extensions python-pptx left undeclared
                ct = b.decode()
                for ext, typ in (('jpg', 'image/jpeg'), ('jpeg', 'image/jpeg'), ('png', 'image/png'), ('fntdata', 'application/x-fontdata')):
                    if any(m.lower().endswith('.' + ext) for m in names) and ('Extension="%s"' % ext) not in ct:
                        ct = ct.replace('<Default ', '<Default Extension="%s" ContentType="%s"/><Default ' % (ext, typ), 1)
                b = ct.encode()
            zo.writestr(n, b)
        zo.close()

# ---------- shapes ----------
def _noline(sh): sh.line.fill.background()
def _shadowless(sh): sh.shadow.inherit = False
def shape(s, kind, x, y, w, h, fill=None, line=None, lw=1.5, name=None):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
    else: _noline(sh)
    _shadowless(sh)
    if name: sh.name = name
    tf = sh.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return sh
def pill(s, x, y, w, h, c, **k):
    sh = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, c, name='Pill', **k); sh.adjustments[0] = 0.5; return sh
def arch(s, x, y, w, h, c, **k):
    sh = shape(s, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, x, y, w, h, c, name='Arch', **k); sh.adjustments[0] = 0.5; sh.adjustments[1] = 0; return sh
def card(s, x, y, w, h, c, line=None, r=0.18):
    sh = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, c, line=line, name='Card'); sh.adjustments[0] = min(0.5, r / min(w, h)); return sh
def circle(s, x, y, d, c): return shape(s, MSO_SHAPE.OVAL, x, y, d, d, c, name='Circle')
def ring(s, x, y, d, c=None, lw=4.5): return shape(s, MSO_SHAPE.OVAL, x, y, d, d, None, line=c or K['lime'], lw=lw, name='Focus ring')
def line(s, x, y, w, c=None, lw=1.5, h=0, dash=False):
    ln = s.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y + h))
    ln.line.color.rgb = rgb(c or K['ink']); ln.line.width = Pt(lw); ln.name = 'Line'
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE; ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return ln
def arrow(s, x, y, c=None, w=0.28, h=0.4):
    return shape(s, MSO_SHAPE.CHEVRON, x, y, w, h, c or K['lime'], name='Arrow')

def tb(s, x, y, w, h, content, size=13, color=None, bold=None, font=None, align=None, valign='t', bullet=False, space_after=None, line=None, wrap=True):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame
    tf.word_wrap = wrap; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[valign]
    base = dict(size=size, color=color or (K['ink']), font=font or BODY)
    if bold is not None: base['bold'] = bold
    set_paras(tf, content, base=base, bullet=bullet, space_after=space_after, align=align, line=line)
    return t
def label_text(sh, content, size=11, color=None, bold=True, align='c', font=None):
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_paras(tf, content, base=dict(size=size, color=color or K['ink'], bold=bold, font=font or BODY), align=align)
    return sh
def tag(s, x, y, t, c, tc=None, w=None):
    sh = pill(s, x, y, w or (0.35 + len(t) * 0.085), 0.34, c); sh.name = 'Tag ' + t
    return label_text(sh, t, 11, tc or K['ink'])
def btn(s, x, y, t, primary=True, dark=False):
    sh = pill(s, x, y, 0.5 + len(t) * 0.1, 0.44, K['lime'] if primary else None, line=None if primary else (K['fl'] if dark else K['forest']))
    sh.name = 'Button'; return label_text(sh, t, 12, K['ink'] if primary else (K['white'] if dark else K['forest']))
def numdot(s, x, y, n, c, tc=None, d=0.5):
    sh = circle(s, x, y, d, c); sh.name = 'Number dot'; return label_text(sh, [Run(str(n), font=HEAD)], 14, tc or K['ink'], bold=False)
def journey(s, x, y, w, colors, d=0.42, lc=None):
    line(s, x + d / 2, y + d / 2, w - d, lc or K['ink'], 2)
    for i, c in enumerate(colors):
        sh = shape(s, MSO_SHAPE.OVAL, x + i * (w - d) / (len(colors) - 1), y, d, d, c, line='FFFFFF', lw=2.5, name='Journey dot')
def trio(s, x, base, sc=1, dark=False):
    arch(s, x, base - 2.2 * sc, 0.9 * sc, 2.2 * sc, K['sky']); arch(s, x + 1.05 * sc, base - 3.1 * sc, 0.9 * sc, 3.1 * sc, K['fl'] if dark else K['forest']); arch(s, x + 2.1 * sc, base - 4.2 * sc, 0.9 * sc, 4.2 * sc, K['lime'])

# ---------- images ----------
GEOM = {'arch': ('round2SameRect', [('adj1', 50000), ('adj2', 0)]), 'pill': ('roundRect', [('adj', 50000)]), 'round': ('roundRect', [('adj', 7000)]),
        'circle': ('ellipse', []), 'rect': ('rect', [])}
def _geom(sp, kind):
    prst, adj = GEOM[kind]
    spPr = sp.find(qn('p:spPr'))
    for g in spPr.findall(qn('a:prstGeom')): spPr.remove(g)
    g = etree.SubElement(spPr, qn('a:prstGeom')); g.set('prst', prst); av = etree.SubElement(g, qn('a:avLst'))
    for n, v in adj:
        gd = etree.SubElement(av, qn('a:gd')); gd.set('name', n); gd.set('fmla', 'val %d' % v)
    # prstGeom must follow xfrm
    xfrm = spPr.find(qn('a:xfrm')); spPr.remove(g); xfrm.addnext(g)
def image(s, path, x, y, w, h, kind='rect', focus=(0.5, 0.5)):
    """Picture cropped to fill the box (no distortion), clipped to a brand shape."""
    im = Image.open(path); iw, ih = im.size; ir, br = iw / ih, w / h
    if path.lower().endswith('.png') and im.mode in ('RGB', 'P', 'L') and os.path.getsize(path) > 300_000:
        jp = os.path.splitext(path)[0] + '.web.jpg'  # big opaque PNG photos -> JPEG to keep the deck small
        if not os.path.exists(jp): im.convert('RGB').save(jp, quality=86)
        path = jp
    pic = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if ir > br:  # too wide: crop sides
        keep = br / ir; cut = 1 - keep; l = cut * focus[0]; pic.crop_left, pic.crop_right = l, cut - l
    else:
        keep = ir / br; cut = 1 - keep; t = cut * focus[1]; pic.crop_top, pic.crop_bottom = t, cut - t
    _geom(pic._element, kind); pic.name = 'Photo'
    return pic
def slot(deck, s, x, y, w, h, kind, what, dark=None):
    """A brand-shaped placeholder for an image or logo the team must add by hand (Format Shape > Fill > Picture keeps the shape)."""
    dark = s._am_dark if dark is None else dark
    k = {'arch': MSO_SHAPE.ROUND_2_SAME_RECTANGLE, 'pill': MSO_SHAPE.ROUNDED_RECTANGLE, 'round': MSO_SHAPE.ROUNDED_RECTANGLE, 'circle': MSO_SHAPE.OVAL, 'rect': MSO_SHAPE.RECTANGLE}[kind]
    sh = shape(s, k, x, y, w, h, K['grid'] if dark else K['ctx'], name='ADD IMAGE: ' + what)
    if kind == 'arch': sh.adjustments[0] = 0.5; sh.adjustments[1] = 0
    if kind == 'pill': sh.adjustments[0] = 0.5
    if kind == 'round': sh.adjustments[0] = 0.07
    label_text(sh, ['ADD: ' + what], 9, K['mutedB'] if dark else K['muted'], bold=True)
    deck.flag(s, what)
    return sh

# ---------- tables ----------
NOSTYLE = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
def _cell_border(cell, side, color, w=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {'b': 'a:lnB', 't': 'a:lnT', 'l': 'a:lnL', 'r': 'a:lnR'}[side]
    for e in tcPr.findall(qn(tag)): tcPr.remove(e)
    ln = etree.SubElement(tcPr, qn(tag)); ln.set('w', str(int(w * 12700)))
    if color:
        sf = etree.SubElement(ln, qn('a:solidFill')); c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', color)
    else: etree.SubElement(ln, qn('a:noFill'))
def table(s, x, y, w, rows, widths=None, size=10, dark=False, head=True, first_bold=False, row_h=0.32, head_h=None):
    nr, nc = len(rows), len(rows[0])
    gt = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr)); t = gt.table
    sid = t._tbl.tblPr.find(qn('a:tableStyleId'))
    if sid is None: sid = etree.SubElement(t._tbl.tblPr, qn('a:tableStyleId'))
    sid.text = NOSTYLE
    t.first_row = False; t.horz_banding = False
    if widths:
        tot = sum(widths)
        for i, cw in enumerate(widths): t.columns[i].width = Inches(w * cw / tot)
    fg = K['white'] if dark else K['ink']; ln = K['grid'] if dark else K['line']
    for r in range(nr):
        t.rows[r].height = Inches(head_h if (head and r == 0 and head_h) else row_h)
        for c in range(nc):
            cell = t.cell(r, c); cell.margin_left = cell.margin_right = Inches(0.08); cell.margin_top = cell.margin_bottom = Inches(0.04)
            ishead = head and r == 0
            if ishead: cell.fill.solid(); cell.fill.fore_color.rgb = rgb(K['panel'] if dark else K['sand'])
            else: cell.fill.background()
            set_paras(cell.text_frame, rows[r][c] if rows[r][c] is not None else '', base=dict(size=size, color=fg, font=BODY, bold=True if (ishead or (first_bold and c == 0)) else False))
            for side in 'tlr': _cell_border(cell, side, None)
            _cell_border(cell, 'b', ln)
    return gt
