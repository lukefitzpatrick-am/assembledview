#!/usr/bin/env python3
"""Brand, overflow and number checks for a deck built from content.json.

Prints OK, or numbered issues, in at most 1500 characters.
"""
import json
import os
import re
import sys
import zipfile

from PIL import ImageFont
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
FONT_DIR = os.path.normpath(os.path.join(HERE, "..", "fonts"))
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
PALETTE = {
    "0F1D13", "FFFFFF", "EFE9DF", "B5D337", "49C7EB", "246646", "4E8F6A",
    "1A2620", "CFC8BA", "E4DED2", "6B6A5E", "AEB9B0", "C9D3C8", "3B4A3F", "2A3A2F",
}
TOKEN = re.compile(r"\$?\d[\d,]*(?:\.\d+)?%?")
EMU = 914400
CHECK = re.compile(r"CHECK (slides\.\d+)")
LOOSE_STOP = re.compile(r"\s[.!?]$")
NUMBER_DOT_NAME = "Number dot"
CONTACT_PLACEHOLDER_IDX = "102"
CHART_PART = re.compile(r"^ppt/(?:charts|embeddings)/")


def font_path(name, italic, bold):
    face = name or ""
    if "Instrument Serif" in face or italic:
        file_name = "InstrumentSerif-Italic.ttf" if italic else "InstrumentSerif-Regular.ttf"
    elif "ExtraBold" in face:
        file_name = "PlusJakartaSans-ExtraBold.ttf"
    elif bold or "Bold" in face:
        file_name = "PlusJakartaSans-Bold.ttf"
    else:
        file_name = "PlusJakartaSans-Regular.ttf"
    return os.path.join(FONT_DIR, file_name)


_FONT = {}


def load_font(name, size_pt, italic, bold):
    px = max(1, int(round(size_pt * 96 / 72)))
    key = (font_path(name, italic, bold), px)
    if key not in _FONT:
        _FONT[key] = ImageFont.truetype(key[0], px)
    return _FONT[key]


def overflows(text_frame, width_in, height_in):
    """True when the frame's text will not fit the box, estimated with PIL."""
    if width_in <= 0.05 or height_in <= 0.05:
        return False
    width_px = width_in * 96
    height_px = height_in * 96
    wrap = text_frame.word_wrap is not False
    used_h = 0
    saw = False
    for paragraph in text_frame.paragraphs:
        runs = [(run.text or "", run) for run in paragraph.runs]
        if not any(text for text, _run in runs):
            continue
        saw = True
        line_w = 0
        line_size = 12
        lines = 1
        too_wide = False
        for text, run in runs:
            size = run.font.size.pt if run.font.size is not None else 12
            font = load_font(run.font.name, size, bool(run.font.italic), bool(run.font.bold))
            for piece in re.findall(r"\S+|\s+", text):
                piece_w = font.getlength(piece)
                if piece.isspace():
                    if line_w > 0 and wrap and line_w + piece_w > width_px:
                        continue
                    line_w += piece_w
                    continue
                if piece_w > width_px + 1:
                    too_wide = True
                if wrap and line_w > 0 and line_w + piece_w > width_px + 1:
                    lines += 1
                    line_w = piece_w
                    line_size = size
                else:
                    line_w += piece_w
                    line_size = max(line_size, size)
        if not wrap and line_w > width_px + 1:
            too_wide = True
        used_h += lines * line_size * 1.15 * 96 / 72
        if too_wide:
            return True
    if not saw:
        return False
    return used_h > height_px + 2


def frame_box(shape):
    width = shape.width / EMU
    height = shape.height / EMU
    if not getattr(shape, "has_text_frame", False):
        return width, height
    frame = shape.text_frame
    width -= ((frame.margin_left or 0) + (frame.margin_right or 0)) / EMU
    height -= ((frame.margin_top or 0) + (frame.margin_bottom or 0)) / EMU
    return width, height


def local(tag):
    return tag.split("}")[-1] if "}" in tag else tag


def qn_tag(name):
    return "{%s}%s" % (NS["p"], name)


def shape_name(element):
    for node in element.iter():
        if local(node.tag) == "cNvPr" and node.get("name"):
            return node.get("name")
    return ""


def texts_of(element):
    return " ".join(node.text or "" for node in element.iter() if local(node.tag) == "t")


def figure_part_skipped(name):
    """Chart workbooks and chart XML are data, not prose figures."""
    return CHART_PART.match(str(name or "").lstrip("/")) is not None


def is_chart_frame(element):
    return any(local(node.tag) == "chart" for node in element.iter())


def canon(raw):
    stripped = raw.replace("$", "").replace(",", "").replace("%", "").replace(" ", "")
    whole, dot, frac = stripped.partition(".")
    whole = whole.lstrip("0") or "0"
    return whole if not dot else f"{whole}.{frac}"


def payload_tokens(json_text):
    return {canon(match) for match in TOKEN.findall(json_text)}


def theme_fonts(zip_file):
    fonts = set()
    for name in zip_file.namelist():
        if not re.match(r"ppt/theme/theme\d+\.xml$", name):
            continue
        xml = zip_file.read(name).decode("utf-8", "replace")
        fonts.update(re.findall(r'typeface="([^"]*)"', xml))
    return fonts


def font_allowed(name, theme):
    if not name or name.startswith("+"):
        return True
    if name == "Instrument Serif" or name.startswith("Plus Jakarta Sans"):
        return True
    return name in theme


def title_of(xml):
    shapes = re.findall(r"<p:sp[\s>][\s\S]*?</p:sp>", xml)
    for shape in shapes:
        if re.search(r'<p:ph[^>]*type="(?:title|ctrTitle)"', shape):
            text = "".join(re.findall(r"<a:t[^>]*>([^<]*)</a:t>", shape)).strip()
            if text:
                return text
    best = ""
    best_size = -1
    for shape in shapes:
        if 'type="sldNum"' in shape:
            continue
        text = "".join(re.findall(r"<a:t[^>]*>([^<]*)</a:t>", shape)).strip()
        if not re.search(r"[A-Za-z]", text):
            continue
        sizes = [int(n) for n in re.findall(r'\ssz="(\d+)"', shape)]
        size = max(sizes) if sizes else 0
        if size >= best_size:
            best = text
            best_size = size
    return best


def content_path(name, slide_path):
    if name and name.startswith("slides."):
        return name
    return slide_path


def placeholder_idx(element):
    for node in element.iter():
        if local(node.tag) == "ph":
            return node.get("idx")
    return None


def slide_part(slide, index):
    part = getattr(getattr(slide, "part", None), "partname", None)
    if part:
        return str(part).lstrip("/")
    return f"ppt/slides/slide{index + 1}.xml"


def stray_figures(deck, allowed):
    """Numbers on the deck that are not in the payload.

    Skips text inside shapes named Number dot, and every text frame on the
    cloned thank-you slide's contact placeholder (body idx 102).
    """
    stray = []
    for index, slide in enumerate(deck.slides):
        number = index + 1
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text or ""
        match = CHECK.search(notes)
        slide_path = match.group(1) if match else f"slides.{index}"
        thank_you = "thank-you" in notes
        tree = slide._element.find(qn_tag("cSld"))
        tree = tree.find(qn_tag("spTree")) if tree is not None else None
        nodes = list(tree) if tree is not None else []
        for element in nodes:
            if local(element.tag) not in ("sp", "graphicFrame"):
                continue
            name = shape_name(element)
            if name == NUMBER_DOT_NAME:
                continue
            if thank_you and placeholder_idx(element) == CONTACT_PLACEHOLDER_IDX:
                continue
            if local(element.tag) == "graphicFrame" and is_chart_frame(element):
                continue
            where = content_path(name, slide_path)
            part = slide_part(slide, index)
            for raw in TOKEN.findall(texts_of(element)):
                normalised = canon(raw)
                if normalised not in allowed:
                    stray.append({
                        "slide": part,
                        "number": number,
                        "where": where,
                        "raw": raw,
                        "normalised": normalised,
                    })
    return stray


def number_issues(pptx_path, json_path):
    with open(json_path, encoding="utf-8") as handle:
        allowed = payload_tokens(handle.read())
    deck = Presentation(pptx_path)
    figures = [
        item for item in stray_figures(deck, allowed)
        if not figure_part_skipped(item["slide"])
    ]
    return {
        "stray": [
            {"slide": item["slide"], "raw": item["raw"], "normalised": item["normalised"]}
            for item in figures
        ],
        "slideCount": len(deck.slides),
    }


def loose_stop(text):
    return " ." in text or LOOSE_STOP.search(text) is not None


def note_loose_stops(issues, number, path, text_frame):
    for paragraph in text_frame.paragraphs:
        text = "".join(run.text or "" for run in paragraph.runs)
        if loose_stop(text):
            issues.append(f"slide {number} {path} loose stop")


def check(pptx_path, json_path):
    issues = []
    with open(json_path, encoding="utf-8") as handle:
        allowed = payload_tokens(handle.read())
    deck = Presentation(pptx_path)
    with zipfile.ZipFile(pptx_path) as zip_file:
        theme = theme_fonts(zip_file)
        slide_parts = sorted(
            (name for name in zip_file.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", name)),
            key=lambda name: int(re.search(r"(\d+)", name).group(1)),
        )
        xmls = [zip_file.read(name).decode("utf-8", "replace") for name in slide_parts]
    stray_by_slide = {}
    for item in stray_figures(deck, allowed):
        stray_by_slide.setdefault(item["number"], []).append(item)

    for index, slide in enumerate(deck.slides):
        number = index + 1
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text or ""
        match = CHECK.search(notes)
        slide_path = match.group(1) if match else f"slides.{index}"
        xml = xmls[index] if index < len(xmls) else ""

        if "<a:gradFill" in xml or "<a:gradFill>" in xml:
            issues.append(f"slide {number} {slide_path} gradFill")
        for colour in re.findall(r"<a:srgbClr[^>]*\sval=\"([0-9A-Fa-f]{6})\"", xml):
            if colour.upper() not in PALETTE:
                issues.append(f"slide {number} {slide_path} srgbClr {colour.upper()}")
        for face in re.findall(r'typeface="([^"]+)"', xml):
            if not font_allowed(face, theme):
                issues.append(f"slide {number} {slide_path} font {face}")
        for hundredths in re.findall(r'\ssz="(\d+)"', xml):
            if int(hundredths) < 800:
                issues.append(f"slide {number} {slide_path} text {int(hundredths) / 100:g}pt")
        title = title_of(xml)
        if title and not title.endswith("."):
            issues.append(f"slide {number} {slide_path}.title no full stop")

        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                base = shape.name[:-6] if shape.name.endswith(".table") else slide_path
                for row_index, row in enumerate(shape.table.rows):
                    row_h = (row.height or 0) / EMU
                    for col_index, cell in enumerate(row.cells):
                        path = f"{base}.columns.{col_index}" if row_index == 0 else f"{base}.rows.{row_index - 1}.{col_index}"
                        col_w = (shape.table.columns[col_index].width or 0) / EMU
                        if overflows(cell.text_frame, col_w - 0.16, row_h - 0.08):
                            issues.append(f"slide {number} {path} overflow")
                        note_loose_stops(issues, number, path, cell.text_frame)
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            if shape_name(shape._element) == "" and not shape.text_frame.text.strip():
                continue
            name = shape.name or ""
            note_loose_stops(issues, number, content_path(name, slide_path), shape.text_frame)
            if "Number dot" in name:
                continue
            box_w, box_h = frame_box(shape)
            if shape.text_frame.text.strip() and overflows(shape.text_frame, box_w, box_h):
                issues.append(f"slide {number} {content_path(name, slide_path)} overflow")

        for placeholder in slide.placeholders:
            if placeholder.placeholder_format.idx is not None and placeholder.placeholder_format.idx >= 4294967295:
                continue
            kind = str(placeholder.placeholder_format.type)
            if "SLIDE_NUMBER" in kind:
                continue
            has_text = placeholder.has_text_frame and placeholder.text_frame.text.strip()
            is_picture = "PICTURE" in kind or (placeholder.shape_type is not None and "PICTURE" in str(placeholder.shape_type))
            if is_picture or not has_text:
                if not has_text:
                    issues.append(f"slide {number} {slide_path}.ph.{placeholder.placeholder_format.idx} empty")

        for item in stray_by_slide.get(number, []):
            issues.append(f"slide {number} {item['where']} number {item['raw']}")

    seen = set()
    unique = []
    for issue in issues:
        if issue not in seen:
            seen.add(issue)
            unique.append(issue)
    return unique


def render(issues):
    if not issues:
        return "OK"
    lines = [f"{number}. {issue}" for number, issue in enumerate(issues, 1)]
    text = "\n".join(lines)
    if len(text) <= 1500:
        return text
    cut = text[:1490]
    end = cut.rfind("\n")
    if end > 0:
        cut = cut[:end]
    return cut + "\n…"


def main():
    if len(sys.argv) >= 2 and sys.argv[1] == "--numbers":
        if len(sys.argv) != 4:
            sys.stderr.write("usage: check_deck.py --numbers deck.pptx report-data.json\n")
            return 2
        sys.stdout.write(json.dumps(number_issues(sys.argv[2], sys.argv[3])) + "\n")
        return 0
    if len(sys.argv) != 4:
        sys.stderr.write("usage: check_deck.py deck.pptx report-data.json content.json\n")
        return 2
    if not os.path.isfile(sys.argv[3]):
        sys.stderr.write("missing content.json\n")
        return 1
    issues = check(sys.argv[1], sys.argv[2])
    sys.stdout.write(render(issues) + "\n")
    return 0 if not issues else 1


if __name__ == "__main__":
    sys.exit(main())
