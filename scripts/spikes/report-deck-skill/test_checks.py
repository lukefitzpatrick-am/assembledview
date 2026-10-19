"""Local checks for the v5 deck builder. No API calls."""
import importlib.util
import json
import os
import re
import sys
import zipfile

from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(HERE, "av-campaign-report-v5", "scripts")
TEMPLATE = os.path.join(HERE, "kit", "assembled-brand-kit", "templates", "Assembled-Media-Template-v5.potx")
DATA = os.path.join(HERE, "out", "hema007-v5-report-data.json")
OUT = os.path.join(HERE, "out")


def load(name):
    path = os.path.join(SCRIPTS, name)
    spec = importlib.util.spec_from_file_location(name.replace(".py", ""), path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def fail(message):
    raise SystemExit(message)


def test_missing_key():
    build = load("build_report.py")
    try:
        build.resolve({"summary": {"spend": "$1.00"}}, "summary.missing", "slides.0.spend")
    except build.BuildError as error:
        if "missing key summary.missing" not in str(error):
            fail(f"missing key message was {error}")
        return
    fail("missing key did not fail")


def test_content_shape_below_6_7_fails():
    am = load("am_rebrand.py")
    check = load("check_deck.py")
    deck = am.Deck(TEMPLATE)
    cover = deck.clone(3)
    deck.notes(cover, "CHECK slides.0 cover sand")
    slide = deck.add("Title Only - White")
    deck.fill(slide, 100, "A short title.")
    deck.notes(slide, "CHECK slides.1 summary stat-row")
    am.card(slide, 0.6, 6.2, 4.0, 0.8, "EFE9DF")
    path = os.path.join(OUT, "below-footer.pptx")
    deck.save(path)
    issues = check.check(path, DATA)
    flagged = [issue for issue in issues if "below 6.7" in issue]
    if flagged != ["slide 2 Card below 6.7"]:
        fail(f"low card was not the only shape below 6.7: {flagged}")


def test_overflow_title():
    am = load("am_rebrand.py")
    check = load("check_deck.py")
    deck = am.Deck(TEMPLATE)
    slide = deck.add("Title Only - White")
    deck.fill(slide, 100, ("A" * 299) + ".")
    path = os.path.join(OUT, "overflow-title.pptx")
    deck.save(path)
    issues = check.check(path, DATA)
    if not any("overflow" in issue and "title" in issue or "overflow" in issue for issue in issues):
        fail(f"300-character title was not flagged: {issues}")


def test_palette():
    am = load("am_rebrand.py")
    check = load("check_deck.py")
    deck = am.Deck(TEMPLATE)
    slide = deck.add("Title Only - White")
    deck.fill(slide, 100, "A short title.")
    am.shape(slide, MSO_SHAPE.RECTANGLE, 1, 2, 1, 1, "00FFFF")
    path = os.path.join(OUT, "palette-flag.pptx")
    deck.save(path)
    issues = check.check(path, DATA)
    if not any("00FFFF" in issue for issue in issues):
        fail(f"00FFFF was not flagged: {issues}")


def test_validate():
    build = load("build_report.py")
    short_list = {
        "slides": [{
            "type": "not-live",
            "variant": "list",
            "title": "Search is {i:\"yet to launch\"}.",
            "rows": [{"label": "Search", "budget": "notLive.0.plannedBudget", "note": "Launch date to be confirmed."}],
        }],
    }
    errors = build.validate_content(short_list)
    if not any("use cards" in error for error in errors):
        fail(f"short list was allowed: {errors}")
    missing_insight = {
        "slides": [{
            "type": "kpi",
            "variant": "table",
            "title": "Rates {i:\"sit below\"} target.",
            "rows": [{"label": "kpis.0.label", "actual": "kpis.0.actualDisplay", "target": "kpis.0.targetDisplay"}],
        }],
    }
    errors = build.validate_content(missing_insight)
    if not any("insight" in error for error in errors):
        fail(f"missing insight was allowed: {errors}")
    bad_variant = {
        "slides": [{
            "type": "summary",
            "variant": "grid",
            "title": "Spend sits {i:\"behind pace\"}.",
        }],
    }
    errors = build.validate_content(bad_variant)
    if not any("variant" in error for error in errors):
        fail(f"bad variant was allowed: {errors}")
    typed = {
        "slides": [{
            "type": "summary",
            "variant": "stat-row",
            "title": "Spend sits {i:\"behind pace\"}.",
            "stats": [{"label": "Spend pace", "key": "35%"}],
            "outcome": "Search is still on the plan.",
        }],
    }
    errors = build.validate_content(typed)
    if not any("value key" in error for error in errors):
        fail(f"typed percent was allowed: {errors}")
    accepted = {
        "slides": [
            {
                "type": "breaker",
                "variant": "black",
                "title": "Where the month {i:\"stands\"}.",
            },
            {
                "type": "recommendations",
                "variant": "cards",
                "title": "What to {i:\"do\"} next.",
                "cards": [
                    {"heading": "Hold the Meta mix.", "body": "Keep the current mix and watch the rate."},
                    {"heading": "Bring Search forward.", "body": "Shift weight toward the terms that convert."},
                    {"heading": "Launch TikTok.", "body": "The planned flight is still waiting to start."},
                ],
                "outcome": "Keep the mix and bring Search forward.",
            },
        ],
    }
    errors = build.validate_content(accepted)
    if errors:
        fail(f"new slide types were rejected: {errors}")
    long_title = {
        "slides": [{
            "type": "breaker",
            "variant": "black",
            "title": ("A" * 60) + ".",
        }],
    }
    errors = build.validate_content(long_title)
    if not any("limit 60" in error for error in errors):
        fail(f"long breaker title was allowed: {errors}")
    long_heading = {
        "slides": [{
            "type": "recommendations",
            "variant": "cards",
            "title": "What to {i:\"do\"} next.",
            "cards": [
                {"heading": "A" * 41, "body": "Keep the current mix."},
                {"heading": "Bring Search forward.", "body": "B" * 221},
            ],
        }],
    }
    errors = build.validate_content(long_heading)
    if not any("limit 40" in error for error in errors):
        fail(f"long heading was allowed: {errors}")
    if not any("limit 220" in error for error in errors):
        fail(f"long body was allowed: {errors}")
    four = {
        "slides": [{
            "type": "recommendations",
            "variant": "cards",
            "title": "What to {i:\"do\"} next.",
            "cards": [
                {"heading": "One.", "body": "First step."},
                {"heading": "Two.", "body": "Second step."},
                {"heading": "Three.", "body": "Third step."},
                {"heading": "Four.", "body": "Fourth step."},
            ],
        }],
    }
    errors = build.validate_content(four)
    if not any("2 to 3" in error for error in errors):
        fail(f"four recommendations were allowed: {errors}")


def slide_text(path):
    text = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                text.append(archive.read(name).decode("utf-8"))
    return "\n".join(text)


def test_pace_and_fallback():
    build = load("build_report.py")
    photos = [
        os.path.join(HERE, "kit", "assembled-brand-kit", "library", name)
        for name in ("pexels-6299928.jpg", "pexels-6786367.jpg", "pexels-13935909.jpg")
    ]
    content = {
        "slides": [
            {
                "type": "summary",
                "variant": "stat-row",
                "title": "Spend sits {i:\"behind pace\"}.",
                "stats": [{"label": "Spend pace", "key": "summary.metrics.spendPacePct"}],
                "outcome": "Search is still on the plan.",
            },
            {
                "type": "not-live",
                "variant": "list",
                "title": "Search is {i:\"yet to launch\"}.",
                "rows": [{
                    "label": "Search",
                    "budget": "notLive.0.plannedBudget",
                    "note": "Launch date to be confirmed.",
                }],
            },
        ],
    }
    with open(DATA, encoding="utf-8") as handle:
        data = json.load(handle)
    out = os.path.join(OUT, "pace-fallback.pptx")
    build.build(content, data, TEMPLATE, out, photos)
    text = slide_text(out)
    if "of expected to date" not in text:
        fail("spend pace label was not extended")
    if "Launch date to be confirmed." not in text:
        fail("short list note was dropped")
    if "Planned" not in text:
        fail("short list did not fall back to cards")


def number_raws(issues):
    return [issue.rsplit(" number ", 1)[-1] for issue in issues if " number " in issue]


def test_number_exemptions():
    """Number dots and the thank-you contact placeholder are the only skips."""
    am = load("am_rebrand.py")
    check = load("check_deck.py")
    deck = am.Deck(TEMPLATE)
    slide = deck.add("Title Only - White")
    deck.fill(slide, 100, "A short title.")
    deck.notes(slide, "CHECK slides.0 summary stat-row")
    box = am.card(slide, 0.8, 2.2, 3.2, 1.1, "FFFFFF")
    am.label_text(box, "$9,999", 18, am.K["ink"])
    am.numdot(slide, 0.8, 3.6, 4, am.K["forest"], am.K["white"], 0.36)
    thank = deck.clone(21)
    deck.notes(thank, "CHECK slides.1 close thank-you")
    deck.fill(thank, 101, "Call extension 42.", base=dict(size=32, color=am.K["white"], font=am.HEAD))
    path = os.path.join(OUT, "number-exemptions.pptx")
    deck.save(path)
    data = os.path.join(OUT, "number-exemptions.json")
    with open(data, "w", encoding="utf-8") as handle:
        json.dump({"summary": {"spend": "$1.00"}}, handle)
    raws = number_raws(check.check(path, data))
    if "$9,999" not in raws:
        fail(f"card $9,999 was not flagged: {raws}")
    if "42" not in raws:
        fail(f"thank-you title 42 was not flagged: {raws}")
    for exempt in ("4", "1", "2", "71", "3205", "97", "2008"):
        if exempt in raws:
            fail(f"{exempt} should be exempt: {raws}")
    report = check.number_issues(path, data)
    report_raws = [item["raw"] for item in report["stray"]]
    if "$9,999" not in report_raws:
        fail(f"shared number check missed $9,999: {report_raws}")
    if "42" not in report_raws:
        fail(f"shared number check missed 42: {report_raws}")
    for exempt in ("4", "1", "2", "71", "3205", "97", "2008"):
        if exempt in report_raws:
            fail(f"shared number check kept {exempt}: {report_raws}")


def test_breaker_and_recommendations_build():
    build = load("build_report.py")
    check = load("check_deck.py")
    photos = [
        os.path.join(HERE, "kit", "assembled-brand-kit", "library", name)
        for name in ("pexels-6299928.jpg", "pexels-6786367.jpg", "pexels-13935909.jpg")
    ]
    with open(DATA, encoding="utf-8") as handle:
        data = json.load(handle)
    shared = {
        "type": "breaker",
        "variant": "black",
        "title": "Where the month {i:\"stands\"}.",
    }
    close = {"type": "close", "variant": "thank-you", "title": "Thanks for {i:\"reading\"}."}

    def cards(count):
        rows = []
        for index in range(count):
            rows.append({
                "heading": f"Step {['one', 'two', 'three'][index]}.",
                "body": ("Keep the current mix. " * 16)[:220],
            })
        return {
            "type": "recommendations",
            "variant": "cards",
            "title": "What to {i:\"do\"} next.",
            "cards": rows,
            "outcome": "Keep the mix and bring Search forward.",
        }

    for count, name in ((2, "recs-two.pptx"), (3, "recs-three.pptx")):
        content = {"slides": [shared, cards(count), close]}
        out = os.path.join(OUT, name)
        build.build(content, data, TEMPLATE, out, photos)
        issues = check.check(out, DATA)
        if issues:
            fail(f"{name} check failed: {issues}")


def run_text(part):
    return part.t if hasattr(part, "t") else part


def test_title_full_stop_sits_on_the_last_word():
    """The full stop is not its own run after the serif phrase."""
    build = load("build_report.py")
    cases = [
        ('Spend sits {i:"behind pace"}.', "behind pace."),
        ('Social (Meta) is {i:"on track"}.', "on track."),
        ('Where the month {i:"stands"}.', "stands."),
        ('Thanks for {i:"reading"}.', "reading."),
        ('Rates sit {i:"below target"}.', "below target."),
    ]
    for title, serif in cases:
        _visible, parts = build.parse_title(title, "title")
        texts = [run_text(part) for part in parts]
        if any(text.strip() == "." for text in texts):
            fail(f"full stop is its own run for {title!r}: {texts}")
        if serif not in texts:
            fail(f"serif run should be {serif!r}, got {texts}")
        if " ." in "".join(texts):
            fail(f"space before the full stop in {texts}")
    _visible, parts = build.parse_title('What to {i:"do"} next.', "title")
    texts = [run_text(part) for part in parts]
    if texts[-1] != " next.":
        fail(f"words after the serif phrase lost the stop: {texts}")
    _visible, parts = build.parse_title('The Curatif {i:"campaign"} report.', "title")
    texts = [run_text(part) for part in parts]
    joined = "".join(texts)
    if not joined.endswith("report."):
        fail(f"cover title stop moved: {texts}")
    if " ." in joined or any(text.strip() == "." for text in texts):
        fail(f"cover title has a loose stop: {texts}")


NOT_LIVE_LINE = (
    "slide {number} (channel): {name} is not live. "
    "Remove this slide and mention {name} in the summary Outcome instead."
)


def not_live_line(number, name="Social (TikTok)"):
    return NOT_LIVE_LINE.format(number=number, name=name)


def assert_not_live(errors, number, name="Social (TikTok)"):
    expected = not_live_line(number, name)
    if expected not in errors:
        fail(f"expected {expected!r}, got {errors}")
    if any("\n" in error for error in errors):
        fail(f"not-live error is more than one line: {errors}")


def test_not_live_channel_has_no_slide():
    build = load("build_report.py")
    data = {
        "channels": [
            {"group": "social_meta", "label": "Social (Meta)", "notLive": False},
            {"group": "social_tiktok", "label": "Social (TikTok)", "notLive": True},
        ],
        "notLive": [{"group": "social_tiktok", "label": "Social (TikTok)"}],
    }
    own_slide = {
        "slides": [{
            "type": "not-live",
            "variant": "cards",
            "title": "TikTok is {i:\"yet to launch\"}.",
            "cards": [{
                "label": "Social (TikTok)",
                "budget": "notLive.0.plannedBudget",
                "status": "Not yet live.",
            }],
        }],
    }
    errors = build.validate_content(own_slide, data)
    assert_not_live(errors, 1)
    by_index = {
        "slides": [{
            "type": "channel",
            "variant": "photo",
            "title": "TikTok is {i:\"yet to launch\"}.",
            "value": "channels.1.spend",
            "body": "Waiting to start.",
            "photo": 0,
        }],
    }
    errors = build.validate_content(by_index, data)
    assert_not_live(errors, 1)
    by_group = {
        "slides": [{
            "type": "channel",
            "variant": "photo",
            "title": "TikTok is {i:\"yet to launch\"}.",
            "value": "channels.social_tiktok.spend",
            "body": "Waiting to start.",
            "photo": 0,
        }],
    }
    errors = build.validate_content(by_group, data)
    assert_not_live(errors, 1)
    live = {
        "slides": [{
            "type": "channel",
            "variant": "photo",
            "title": "Meta is {i:\"on track\"}.",
            "value": "channels.0.spend",
            "body": "Delivery is in.",
            "photo": 0,
        }],
    }
    errors = build.validate_content(live, data)
    if any("notLive" in error for error in errors):
        fail(f"live channel was rejected: {errors}")
    summary = {
        "slides": [{
            "type": "summary",
            "variant": "stat-row",
            "title": "Spend sits {i:\"behind pace\"}.",
            "stats": [{"label": "Spend", "key": "summary.spend"}],
            "outcome": "TikTok has not started.",
        }],
    }
    errors = build.validate_content(summary, data)
    if any("not live" in error for error in errors):
        fail(f"summary outcome line was rejected: {errors}")


def test_not_live_allows_summary_kpi_and_rates():
    """a to d: only a slide whose own channel is not live fails."""
    import shutil
    import subprocess
    import tempfile

    build = load("build_report.py")
    data = {
        "channels": [
            {"group": "social_meta", "label": "Social (Meta)", "notLive": False},
            {"group": "social_tiktok", "label": "Social (TikTok)", "notLive": True},
        ],
        "notLive": [{"group": "social_tiktok", "label": "Social (TikTok)"}],
    }
    base_path = os.path.join(OUT, "ua5-content.json")
    if not os.path.isfile(base_path):
        fail(f"missing hand-written content {base_path}")
    with open(base_path, encoding="utf-8") as handle:
        base = json.load(handle)
    for slide in base["slides"]:
        if slide.get("type") == "recommendations" and "outcome" not in slide:
            slide["outcome"] = "Keep Meta and push Search."

    def without_not_live(errors):
        return [error for error in errors if "is not live" in error]

    hand_written = build.validate_content(base, data)
    if without_not_live(hand_written):
        fail(f"hand-written Curatif content was rejected: {hand_written}")

    outcome_only = json.loads(json.dumps(base))
    for slide in outcome_only["slides"]:
        if slide.get("type") == "summary":
            slide["outcome"] = "Meta and Search are live. TikTok has not started."
    outcome_errors = build.validate_content(outcome_only, data)
    if without_not_live(outcome_errors):
        fail(f"summary Outcome mention was rejected: {outcome_errors}")

    kpi = json.loads(json.dumps(base))
    for slide in kpi["slides"]:
        if slide.get("type") == "kpi":
            slide["rows"].append({
                "label": "channels.1.label",
                "actual": "channels.1.spend",
                "target": "channels.1.plannedBudget",
            })
    kpi_errors = build.validate_content(kpi, data)
    if without_not_live(kpi_errors):
        fail(f"KPI TikTok row was rejected: {kpi_errors}")

    rates = json.loads(json.dumps(base))
    rates["slides"].insert(-2, {
        "type": "appendix",
        "variant": "table",
        "title": "Rates by {i:\"channel\"}.",
        "columns": ["Channel", "CPM"],
        "rows": [["channels.1.label", "channels.1.metrics.cpm"]],
        "insight": "TikTok stays in the rates table.",
    })
    rates_errors = build.validate_content(rates, data)
    if without_not_live(rates_errors):
        fail(f"rates TikTok row was rejected: {rates_errors}")

    channel = json.loads(json.dumps(base))
    channel["slides"].insert(3, {
        "type": "channel",
        "variant": "photo",
        "title": "TikTok is {i:\"yet to launch\"}.",
        "channel": "TikTok",
        "value": "channels.1.spend",
        "body": "Waiting to start.",
        "photo": 0,
    })
    channel_errors = build.validate_content(channel, data)
    assert_not_live(channel_errors, 4)

    script = os.path.join(SCRIPTS, "write_content.py")
    work = tempfile.mkdtemp(prefix="not-live-")
    data_path = os.path.join(work, "report-data.json")
    try:
        with open(data_path, "w", encoding="utf-8") as handle:
            json.dump(data, handle)
        cases = {
            "a": base,
            "b": outcome_only,
            "c": kpi,
            "d": channel,
        }
        for label, content in cases.items():
            proc = subprocess.run(
                [sys.executable, script, data_path],
                input=json.dumps(content),
                text=True,
                capture_output=True,
                cwd=work,
            )
            if label == "d":
                if proc.returncode == 0:
                    fail(f"write_content.py accepted a TikTok channel slide: {proc.stdout}")
                if not_live_line(4) not in proc.stdout:
                    fail(f"write_content.py error format was {proc.stdout!r}")
            elif proc.returncode != 0 or proc.stdout.strip() != "OK":
                fail(f"write_content.py case {label} failed: {proc.stdout}{proc.stderr}")
    finally:
        shutil.rmtree(work, ignore_errors=True)


def test_kpi_value_drops_repeated_unit():
    build = load("build_report.py")
    photos = [
        os.path.join(HERE, "kit", "assembled-brand-kit", "library", name)
        for name in ("pexels-6299928.jpg", "pexels-6786367.jpg", "pexels-13935909.jpg")
    ]
    data = {
        "kpis": [
            {"label": "Frequency", "targetDisplay": "1.0 frequency", "actualDisplay": "Not yet reported"},
            {"label": "CTR", "targetDisplay": "4.00%", "actualDisplay": "2.16%"},
        ],
    }
    content = {
        "slides": [{
            "type": "kpi",
            "variant": "table",
            "title": "Rates sit {i:\"below target\"}.",
            "rows": [
                {"label": "kpis.0.label", "actual": "kpis.0.actualDisplay", "target": "kpis.0.targetDisplay"},
                {"label": "kpis.1.label", "actual": "kpis.1.actualDisplay", "target": "kpis.1.targetDisplay"},
            ],
            "insight": "Click through rate is behind its target.",
        }],
    }
    out = os.path.join(OUT, "kpi-units.pptx")
    build.build(content, data, TEMPLATE, out, photos)
    text = slide_text(out)
    if "1.0 frequency" in text:
        fail("frequency unit was repeated in the value column")
    if ">1.0<" not in text and "1.0" not in text:
        fail("frequency value was dropped")
    if "4.00%" not in text or "2.16%" not in text:
        fail(f"percent KPI values were changed: {text[:500]}")


A_RPR = "{http://schemas.openxmlformats.org/drawingml/2006/main}rPr"


def paint_runs(shape):
    rows = []
    if not getattr(shape, "has_text_frame", False):
        return rows
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            props = run._r.find(A_RPR)
            spc = props.get("spc") if props is not None else None
            rows.append((run.text or "", run.font.name, bool(run.font.italic), spc))
    return rows


def test_stop_attaches_on_sans_and_serif_runs():
    """A closing stop sits on the last run with no space, for a sans ending and a serif ending."""
    build = load("build_report.py")
    check = load("check_deck.py")
    deck = build.Deck(TEMPLATE)
    sans = deck.add("Title Only - White")
    _visible, runs = build.parse_title('What to {i:"do"} next .', "title")
    deck.fill(sans, 100, runs, base=dict(size=36, color=build.K["ink"], font=build.HEAD), align="l")
    build.tb(sans, 0.6, 2.2, 6, 0.6, "Hold the Meta mix.", size=14, color=build.K["ink"], font=build.HEAD)
    build.tb(sans, 0.6, 2.9, 6, 0.6, "Push Search harder.", size=14, color=build.K["ink"], font=build.HEAD)
    serif = deck.add("Title Only - White")
    _visible, runs = build.parse_title('Spend sits {i:"behind pace"}.', "title")
    deck.fill(serif, 100, runs, base=dict(size=36, color=build.K["ink"], font=build.HEAD), align="l")
    path = os.path.join(OUT, "stops.pptx")
    deck.save(path)
    loose = [issue for issue in check.check(path, DATA) if "loose stop" in issue]
    if loose:
        fail(f"attached stops were flagged: {loose}")

    sans_title = paint_runs(deck.ph(sans, 100))
    text, face, italic, spc = sans_title[-1]
    joined = "".join(part[0] for part in sans_title)
    if " ." in joined or text.strip() == "." or not text.endswith("."):
        fail(f"sans title stop is loose: {sans_title}")
    if face != build.HEAD or italic:
        fail(f"sans title stop left the sans run: {sans_title}")
    if spc is None or int(spc) >= 0:
        fail(f"sans title stop was not pulled in: {sans_title}")

    for shape in sans.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        painted = paint_runs(shape)
        heading = "".join(part[0] for part in painted)
        if heading not in ("Hold the Meta mix.", "Push Search harder."):
            continue
        tail = painted[-1]
        if " ." in heading or tail[0].strip() == "." or not tail[0].endswith("."):
            fail(f"heading stop is loose: {painted}")
        if tail[1] != build.HEAD or tail[3] is None or int(tail[3]) >= 0:
            fail(f"heading stop was not pulled onto the sans run: {painted}")

    serif_title = paint_runs(deck.ph(serif, 100))
    text, face, italic, spc = serif_title[-1]
    joined = "".join(part[0] for part in serif_title)
    if " ." in joined or text != "behind pace." or not italic or face != "Instrument Serif":
        fail(f"serif title stop left the serif run: {serif_title}")
    if spc not in (None, "0"):
        fail(f"serif stop was retargeted: {serif_title}")

    bad = build.Deck(TEMPLATE)
    slide = bad.add("Title Only - White")
    frame = bad.ph(slide, 100).text_frame
    frame.clear()
    frame.paragraphs[0].text = "What to do next ."
    bad_path = os.path.join(OUT, "stops-loose.pptx")
    bad.save(bad_path)
    flagged = [issue for issue in check.check(bad_path, DATA) if "loose stop" in issue]
    if not flagged:
        fail("a space before the full stop was allowed")


def photos():
    return [
        os.path.join(HERE, "kit", "assembled-brand-kit", "library", name)
        for name in ("pexels-6299928.jpg", "pexels-6786367.jpg", "pexels-13935909.jpg")
    ]


def chart_payload():
    return {
        "summary": {"spend": "$10.00", "clicks": "4"},
        "channels": [{
            "group": "social_meta",
            "label": "Social (Meta)",
            "notLive": False,
            "spend": "$10.00",
            "impressions": "80",
            "clicks": "4",
            "metrics": {"ctr": "5.00%"},
        }],
        "kpis": [{
            "metric": "conversion_rate",
            "label": "Conv. Rate",
            "targetDisplay": "5.00%",
            "actualDisplay": "3.00%",
        }],
        "charts": [
            {
                "scope": "campaign",
                "group": None,
                "label": "All channels",
                "left": {"key": "spend", "label": "Spend", "format": "money"},
                "right": {"key": "clicks", "label": "Clicks", "format": "count"},
                "points": [["2026-07-01", 10, 4], ["2026-07-02", 0, 0]],
                "note": "Reported daily spend. Excludes assumed and estimated spend.",
            },
            {
                "scope": "channel",
                "group": "social_meta",
                "label": "Social (Meta)",
                "left": {"key": "amount_spent", "label": "Spend", "format": "money"},
                "right": {"key": "impressions", "label": "Impressions", "format": "count"},
                "points": [["2026-07-01", 10, 80], ["2026-07-02", 0, 0]],
                "note": None,
            },
        ],
    }


def tick_label_skip(point_count):
    build = load("build_report.py")
    data = chart_payload()
    data["charts"][0]["points"] = [
        [f"2026-01-{(index % 28) + 1:02d}", 0, 0] for index in range(point_count)
    ]
    data["charts"][0]["note"] = None
    content = {"slides": [{
        "type": "chart",
        "variant": "delivery",
        "title": "Spend is {i:\"in\"}.",
        "chart": "charts.0",
        "stats": [{"label": "Spend", "key": "summary.spend"}],
        "insight": "Daily spend is in the chart.",
    }]}
    out = os.path.join(OUT, f"skip-{point_count}.pptx")
    build.build(content, data, TEMPLATE, out, photos())
    with zipfile.ZipFile(out) as archive:
        chart = next(name for name in archive.namelist() if name.startswith("ppt/charts/chart"))
        xml = archive.read(chart).decode("utf-8")
    axis = re.search(r"<c:catAx>.*?</c:catAx>", xml, re.DOTALL)
    if axis is None:
        fail(f"{point_count}-point chart has no category axis")
    return axis.group(0)


def test_date_labels_skip_to_eight():
    month = tick_label_skip(30)
    if 'tickLblSkip val="4"' not in month or 'wrap="none"' not in month:
        fail(f"30-point chart did not skip to 4 one-line labels: {month}")
    campaign = tick_label_skip(105)
    if 'tickLblSkip val="14"' not in campaign or 'wrap="none"' not in campaign:
        fail(f"105-point chart did not skip to 14 one-line labels: {campaign}")


def test_chart_slide_builds_and_opens():
    os.makedirs(OUT, exist_ok=True)
    build = load("build_report.py")
    check = load("check_deck.py")
    data = chart_payload()
    content = {"slides": [{
        "type": "chart",
        "variant": "delivery",
        "title": "Spend is {i:\"in\"}.",
        "chart": "charts.0",
        "stats": [{"label": "Spend", "key": "summary.spend"}],
        "insight": "Daily spend is in the chart.",
    }]}
    data_path = os.path.join(OUT, "chart-payload.json")
    out = os.path.join(OUT, "chart-slide.pptx")
    with open(data_path, "w", encoding="utf-8") as handle:
        json.dump(data, handle)
    build.build(content, data, TEMPLATE, out, photos())
    issues = check.check(out, data_path)
    if issues:
        fail(f"chart slide did not open clean: {issues}")
    report = check.number_issues(out, data_path)
    if report["stray"]:
        fail(f"chart figures were treated as prose: {report['stray']}")
    with zipfile.ZipFile(out) as archive:
        names = archive.namelist()
        charts = [name for name in names if name.startswith("ppt/charts/")]
        if not charts:
            fail("chart slide wrote no chart part")
        xml = archive.read(charts[0]).decode("utf-8")
        if "<c:lineChart" not in xml or xml.count("<c:valAx") < 2:
            fail("combo chart XML is missing the line series or the second axis")
        if 'legendPos val="tr"' not in xml:
            fail("chart legend is not top right")
        if "<c:manualLayout>" not in xml or 'h val="0.08"' not in xml:
            fail("chart legend is not the top 8% of the frame")
        if 'name="smooth" val="0"' not in xml and 'smooth val="0"' not in xml:
            fail("chart line is not set smooth off")
    text = slide_text(out)
    if "Reported daily spend. Excludes assumed and estimated spend." not in text:
        fail(f"chart note was not drawn: {text}")
    from pptx import Presentation
    Presentation(out)


def test_channel_chart_renders_four_fixed_stats():
    os.makedirs(OUT, exist_ok=True)
    build = load("build_report.py")
    check = load("check_deck.py")
    data = chart_payload()
    content = {"slides": [{
        "type": "channel",
        "variant": "chart",
        "title": "Meta is {i:\"on track\"}.",
        "chart": "charts.1",
        "insight": "Meta spend and impressions are in the chart.",
    }]}
    data_path = os.path.join(OUT, "channel-chart-payload.json")
    out = os.path.join(OUT, "channel-chart.pptx")
    with open(data_path, "w", encoding="utf-8") as handle:
        json.dump(data, handle)
    build.build(content, data, TEMPLATE, out, photos())
    issues = check.check(out, data_path)
    if issues:
        fail(f"channel chart slide did not open clean: {issues}")
    text = slide_text(out)
    for label in ("Spend", "Impressions", "Clicks", "CTR", "$10.00", "80", "4", "5.00%"):
        if label not in text:
            fail(f"channel chart is missing fixed stat {label}: {text}")
    missing = {
        "slides": [{
            "type": "channel",
            "variant": "chart",
            "title": "Search is {i:\"on track\"}.",
            "chart": "charts.9",
            "insight": "There is no chart for this channel.",
        }],
    }
    errors = build.validate_content(missing, data)
    if not any(error.endswith("use variant cards") and "\n" not in error for error in errors):
        fail(f"missing channel chart was not sent to cards: {errors}")


def test_recommendations_without_outcome_fail():
    build = load("build_report.py")
    content = {
        "slides": [{
            "type": "recommendations",
            "variant": "cards",
            "title": "What to {i:\"do\"} next.",
            "cards": [
                {"heading": "Hold the mix.", "body": "Keep the current mix."},
                {"heading": "Push Search.", "body": "Move weight toward Search."},
            ],
        }],
    }
    errors = build.validate_content(content)
    outcome = [error for error in errors if "outcome" in error]
    if len(outcome) != 1 or "\n" in outcome[0]:
        fail(f"missing outcome should be one line: {errors}")
    if not outcome[0].startswith("slides.0.outcome:"):
        fail(f"outcome error should name the slide and field: {outcome[0]}")


def detail_payload():
    return {
        "channels": [
            {"group": "social_meta", "label": "Social (Meta)", "notLive": False},
            {"group": "digital_display", "label": "Digital Display", "notLive": False},
            {"group": "social_tiktok", "label": "Social (TikTok)", "notLive": True},
        ],
        "detail": [
            {
                "group": "social_meta",
                "label": "Social (Meta)",
                "level": "ad_group",
                "note": "Prospecting leads this table.",
                "rows": [{
                    "name": "Prospecting",
                    "spend": "$12.00",
                    "impressions": "100",
                    "clicks": "5",
                    "results": "2",
                    "ctr": "5.00%",
                }],
            },
            {
                "group": "digital_display",
                "label": "Digital Display",
                "level": "ad_group",
                "note": "Placements ranked by impressions. Spend is reported at line level.",
                "rows": [{
                    "name": "NewsCorp-Video",
                    "impressions": "249,317",
                    "clicks": "40",
                    "results": "Not yet reported",
                    "ctr": "0.02%",
                }],
            },
        ],
    }


def shape_bottom(shape):
    return (shape.top + shape.height) / 914400


def test_detail_table_slide():
    os.makedirs(OUT, exist_ok=True)
    build = load("build_report.py")
    check = load("check_deck.py")
    data = detail_payload()
    content = {"slides": [
        {
            "type": "detail",
            "variant": "table",
            "title": "Meta ad sets {i:\"lead\"} the spend.",
            "detail": "detail.0",
            "insight": "Prospecting leads the spend.",
        },
        {
            "type": "detail",
            "variant": "table",
            "title": "Display placements are {i:\"ranked\"}.",
            "detail": "detail.1",
            "insight": "NewsCorp leads the placements.",
        },
    ]}
    data_path = os.path.join(OUT, "detail-payload.json")
    out = os.path.join(OUT, "detail-slide.pptx")
    with open(data_path, "w", encoding="utf-8") as handle:
        json.dump(data, handle)
    build.build(content, data, TEMPLATE, out, photos())
    issues = check.check(out, data_path)
    if issues:
        fail(f"detail slides did not open clean: {issues}")
    from pptx import Presentation
    deck = Presentation(out)
    expected = (
        ["Ad set", "Spend", "Impressions", "Clicks", "CTR", "Results"],
        ["Placement", "Impressions", "Clicks", "CTR"],
    )
    for number, slide in enumerate(deck.slides, 1):
        named = {shape.name: shape for shape in slide.shapes}
        table = next((shape for shape in slide.shapes if getattr(shape, "has_table", False)), None)
        note = named.get(f"slides.{number - 1}.note")
        insight = named.get(f"slides.{number - 1}.insight")
        if table is None or note is None or insight is None:
            fail(f"slide {number} is missing the table, note or insight")
        header = [cell.text for cell in table.table.rows[0].cells]
        if header != expected[number - 1]:
            fail(f"slide {number} columns were {header}")
        for shape in (table, note, insight):
            bottom = shape_bottom(shape)
            if bottom > 6.7:
                fail(f"slide {number} {shape.name} ends at {bottom:.2f}, past 6.7")
        if abs(table.table.rows[0].height / 914400 - 0.42) > 0.01:
            fail("detail table row height is not 0.42")
    first = next(shape for shape in deck.slides[0].shapes if getattr(shape, "has_table", False))
    if first.table.rows[1].cells[1].text != "$12.00":
        fail(f"spend cell was {first.table.rows[1].cells[1].text}")
    missing = {"slides": [{
        "type": "detail",
        "variant": "table",
        "title": "This key is {i:\"missing\"}.",
        "detail": "detail.9",
        "insight": "The table is not in the payload.",
    }]}
    errors = build.validate_content(missing, data)
    if len(errors) != 1 or "\n" in errors[0] or not errors[0].startswith("slides.0.detail:"):
        fail(f"missing detail key should be one line naming the field: {errors}")
    if "rows" not in errors[0] or "detail." not in errors[0]:
        fail(f"missing detail key should name the fix: {errors[0]}")
    parked = {"slides": [{
        "type": "detail",
        "variant": "table",
        "title": "TikTok is {i:\"not live\"}.",
        "detail": "detail.0",
        "insight": "This group is not live.",
    }]}
    parked_data = detail_payload()
    parked_data["detail"][0]["group"] = "social_tiktok"
    parked_errors = build.validate_content(parked, parked_data)
    if len(parked_errors) != 1 or not parked_errors[0].startswith("slides.0.detail:"):
        fail(f"a detail group that is not live should be one field error: {parked_errors}")


def test_conversion_rate_note():
    os.makedirs(OUT, exist_ok=True)
    build = load("build_report.py")
    check = load("check_deck.py")
    data = chart_payload()
    content = {"slides": [{
        "type": "kpi",
        "variant": "table",
        "title": "Rates sit {i:\"below target\"}.",
        "rows": [{
            "label": "kpis.0.label",
            "actual": "kpis.0.actualDisplay",
            "target": "kpis.0.targetDisplay",
        }],
        "insight": "Conversion rate is behind its target.",
    }]}
    data_path = os.path.join(OUT, "conversion-payload.json")
    out = os.path.join(OUT, "conversion-note.pptx")
    with open(data_path, "w", encoding="utf-8") as handle:
        json.dump(data, handle)
    build.build(content, data, TEMPLATE, out, photos())
    text = slide_text(out)
    if build.CONVERSION_NOTE not in text:
        fail(f"conversion rate note was not added: {text}")
    issues = check.check(out, data_path)
    if issues:
        fail(f"conversion note slide did not open clean: {issues}")


def main():
    os.makedirs(OUT, exist_ok=True)
    test_missing_key()
    test_overflow_title()
    test_content_shape_below_6_7_fails()
    test_palette()
    test_validate()
    test_pace_and_fallback()
    test_number_exemptions()
    test_breaker_and_recommendations_build()
    test_title_full_stop_sits_on_the_last_word()
    test_not_live_channel_has_no_slide()
    test_not_live_allows_summary_kpi_and_rates()
    test_kpi_value_drops_repeated_unit()
    test_stop_attaches_on_sans_and_serif_runs()
    test_chart_slide_builds_and_opens()
    test_date_labels_skip_to_eight()
    test_channel_chart_renders_four_fixed_stats()
    test_recommendations_without_outcome_fail()
    test_conversion_rate_note()
    test_detail_table_slide()
    print("tests OK")


if __name__ == "__main__":
    main()
