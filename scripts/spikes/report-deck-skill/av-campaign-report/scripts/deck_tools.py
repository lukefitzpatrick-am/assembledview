"""
Helper functions for building Assembled Media decks from the bundled template.

Usage pattern:
    from deck_tools import load_template, duplicate_slide, keep_slides, move_slide, set_text

    prs = load_template()               # opens assets/assembled-template.pptx
    duplicate_slide(prs, 34)            # returns new slide appended at end (0-indexed source)
    keep_slides(prs, [0, 2, 34, 68])    # keeps only these indices, in this order
    prs.save(os.path.join(os.environ['OUTPUT_DIR'], 'deck.pptx'))

Work with 0-indexed slide positions. The slide catalogue in
references/slide-catalogue.md lists slides 1-indexed - subtract 1.

Recommended build order:
1. duplicate_slide() for any slide type needed more than once (duplicates append to the end)
2. Populate text/pictures on every slide you intend to keep
3. keep_slides() last - it also handles final ordering
"""

import copy
import os
from pptx import Presentation

SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE_PATH = os.path.join(SKILL_DIR, 'assets', 'assembled-template.pptx')


def load_template(path=TEMPLATE_PATH):
    return Presentation(path)


def duplicate_slide(prs, index):
    """Duplicate the slide at `index` (0-based). The copy is appended at the end.
    Copies all shapes and non-layout relationships (images, media), remapping
    relationship IDs inside the copied XML so pictures keep working."""
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)

    # Strip placeholders the layout auto-inserted; we copy the source's shapes instead
    for shape in list(dest.shapes):
        shape._element.getparent().remove(shape._element)

    for shape in source.shapes:
        dest.shapes._spTree.append(copy.deepcopy(shape._element))

    # Copy relationships (images, media, hyperlinks), remapping rIds
    rid_map = {}
    for rid, rel in source.part.rels.items():
        if 'slideLayout' in rel.reltype or 'notesSlide' in rel.reltype:
            continue
        if rel.is_external:
            new_rid = dest.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            new_rid = dest.part.rels.get_or_add(rel.reltype, rel._target)
        rid_map[rid] = new_rid

    # Rewrite any r:embed / r:link / r:id attributes in the copied shape tree
    if rid_map:
        ns_r = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
        for el in dest.shapes._spTree.iter():
            for attr in list(el.attrib):
                if attr.startswith('{%s}' % ns_r):
                    old = el.attrib[attr]
                    if old in rid_map:
                        el.attrib[attr] = rid_map[old]
    return dest


def keep_slides(prs, indices):
    """Keep only the slides at the given 0-based indices, reordered to match
    the order of `indices`. Call this LAST, after duplicating and populating."""
    sldIdLst = prs.slides._sldIdLst
    slide_ids = list(sldIdLst)
    keep_set = set(indices)

    # Remove slides not kept
    for i, sid in enumerate(slide_ids):
        if i not in keep_set:
            sldIdLst.remove(sid)

    # Reorder remaining to match indices order
    remaining = {slide_ids[i]: i for i in indices}
    ordered = sorted(sldIdLst, key=lambda sid: indices.index(remaining[sid]))
    for sid in ordered:
        sldIdLst.remove(sid)
    for sid in ordered:
        sldIdLst.append(sid)


def move_slide(prs, from_index, to_index):
    sldIdLst = prs.slides._sldIdLst
    slides = list(sldIdLst)
    sldIdLst.remove(slides[from_index])
    sldIdLst.insert(to_index, slides[from_index])


def set_text(shape, text, keep_format=True):
    """Replace a placeholder's text while preserving the template's run formatting.
    Accepts a string ('\n' creates new paragraphs) or a list of strings (one per paragraph)."""
    tf = shape.text_frame
    paras = text.split('\n') if isinstance(text, str) else list(text)

    # Capture formatting from the first existing run, if any
    src_rPr = None
    if tf.paragraphs and tf.paragraphs[0].runs:
        src_rPr = tf.paragraphs[0].runs[0]._r.find(
            '{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')

    tf.text = paras[0]
    for p in paras[1:]:
        para = tf.add_paragraph()
        para.text = p

    if keep_format and src_rPr is not None:
        for para in tf.paragraphs:
            for run in para.runs:
                existing = run._r.find(
                    '{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
                if existing is not None:
                    run._r.remove(existing)
                run._r.insert(0, copy.deepcopy(src_rPr))


def placeholders(slide):
    """Debug helper - print placeholder idx, type, geometry and current text.
    Call this only on slides already kept. Do not run it across the template.
    Placeholders under ~0.3in are decorative anchors - never put text in them."""
    for shape in slide.shapes:
        if shape.is_placeholder:
            pf = shape.placeholder_format
            txt = shape.text_frame.text[:40] if shape.has_text_frame else ''
            w, h = shape.width / 914400, shape.height / 914400
            tiny = '  << TINY, DO NOT FILL' if (w < 0.3 or h < 0.3) else ''
            print(f"idx={pf.idx:3} type={pf.type} x={shape.left/914400:5.2f} "
                  f"y={shape.top/914400:5.2f} w={w:5.2f} h={h:5.2f} "
                  f"text='{txt}'{tiny}")


def fill_by_idx(slide, mapping):
    """Fill placeholders by their placeholder idx: fill_by_idx(slide, {16: "Item one", 17: "Item two"}).
    This is the ONLY safe way to fill multi-placeholder slides - shape order does
    not match visual order, and some placeholders are tiny decorative anchors."""
    by_idx = {s.placeholder_format.idx: s for s in slide.shapes if s.is_placeholder}
    for idx, text in mapping.items():
        if idx not in by_idx:
            raise KeyError(f"No placeholder with idx={idx} on this slide. "
                           f"Available: {sorted(by_idx)}")
        set_text(by_idx[idx], text)
